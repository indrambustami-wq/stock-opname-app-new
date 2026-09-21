import base64
import io
import re
import warnings
from datetime import datetime
from html import escape as html_escape

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ============================================================
# STREAMLIT CONFIG
# ============================================================

st.set_page_config(
    page_title="Laporan Analisa Stock Opname",
    page_icon="📦",
    layout="wide",
)


# Ukuran font dibuat lebih kecil supaya halaman lebih ringkas:
# - judul section (## KPI, ## Rekapitulasi Kategori, dst)
# - angka & label KPI (st.metric)
# Isi tabel (st.dataframe) digambar di kanvas oleh Streamlit,
# jadi ukuran hurufnya tidak bisa diubah lewat CSS.
st.markdown(
    """
    <style>
    /* ========================================================
       STOCK OPNAME PRO - MOBILE FIRST UI
       ======================================================== */
    :root {
        --so-navy: #1F497D;
        --so-blue: #2F75B5;
        --so-soft: #F4F7FB;
        --so-border: #DCE5F0;
    }

    section.main h1 {
        font-size: 1.55rem !important;
        font-weight: 800 !important;
        letter-spacing: -0.02em;
        margin-bottom: 0.25rem !important;
    }
    section.main h2,
    [data-testid="stMain"] h2 {
        font-size: 1.08rem !important;
        font-weight: 750 !important;
        padding: 0.55rem 0 0.3rem 0 !important;
    }
    [data-testid="stMetric"] {
        background: linear-gradient(180deg, #FFFFFF 0%, #F7FAFE 100%);
        border: 1px solid var(--so-border);
        border-radius: 16px;
        padding: 0.75rem 0.9rem;
        box-shadow: 0 3px 12px rgba(31,73,125,0.07);
        min-height: 92px;
    }
    [data-testid="stMetricValue"],
    [data-testid="stMetricValue"] * {
        font-size: 1.28rem !important;
        line-height: 1.25 !important;
        font-weight: 800 !important;
    }
    [data-testid="stMetricLabel"],
    [data-testid="stMetricLabel"] * {
        font-size: 0.74rem !important;
        font-weight: 650 !important;
    }
    [data-testid="stFileUploader"] {
        border: 1px dashed #AFC4DA;
        border-radius: 16px;
        padding: 0.35rem;
        background: #F8FBFF;
    }
    div.stButton > button,
    div.stDownloadButton > button {
        min-height: 46px;
        border-radius: 12px;
        font-weight: 700;
    }
    [data-testid="stDataFrame"] {
        border-radius: 12px;
        overflow: hidden;
        border: 1px solid var(--so-border);
    }
    .so-hero {
        background: linear-gradient(135deg, #1F497D 0%, #2F75B5 100%);
        color: white;
        border-radius: 20px;
        padding: 1.05rem 1.1rem;
        margin-bottom: 1rem;
        box-shadow: 0 8px 24px rgba(31,73,125,0.20);
    }
    .so-hero-title {
        font-size: 1.35rem;
        font-weight: 800;
        margin: 0;
    }
    .so-hero-sub {
        opacity: .88;
        font-size: .82rem;
        margin-top: .2rem;
    }
    @media (max-width: 768px) {
        section.main h1 { font-size: 1.3rem !important; }
        section.main h2 { font-size: 1rem !important; }
        [data-testid="stMetric"] { min-height: 82px; padding: .65rem .75rem; }
        [data-testid="stMetricValue"], [data-testid="stMetricValue"] * { font-size: 1.05rem !important; }
        [data-testid="stMetricLabel"], [data-testid="stMetricLabel"] * { font-size: .68rem !important; }
        .block-container { padding-left: .7rem !important; padding-right: .7rem !important; }
    }
    </style>
    <div class="so-hero">
      <div class="so-hero-title">📦 Stock Opname Pro</div>
      <div class="so-hero-sub">Analisa • Rekap • Deduct • Laporan</div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# CONSTANTS
# ============================================================

NAVY = "1F497D"
YELLOW = "FFFF00"
WHITE = "FFFFFF"
BLACK = "000000"
RED = "FF0000"

# Format label kolon otomatis (mis. "Shopbag" -> tampil
# "Shopbag                    :"). Titik dua selalu rata ke
# sisi kanan sel, tidak perlu diketik manual di teks.
COLON_LABEL_FORMAT = '@* ":"'
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_GRAY = "F2F2F2"

THIN_GRAY = Side(
    style="thin",
    color="B7B7B7",
)

DOUBLE_BLACK = Side(
    style="double",
    color=BLACK,
)

SHOPBAG_KEYWORDS = [
    "Shopping Bag",
    "Laundry",
    "Paper Bag",
    "Bag",
    "LB",
    "PBB",
    "SB",
]

CATEGORIES = [
    "KLOP",
    "SHOPBAG",
    "TERTUKAR",
    "MINUS",
    "PLUS",
]

# ============================================================
# MASTER LOKASI
# ============================================================
# Sumber: MASTER_LOKASI.xlsx (Location Kode -> Initial).
# Dipakai untuk otomatis menampilkan kode Initial di
# belakang Lokasi / Loc, mis. "HWR0019 ( HGS )".
# Kalau ada lokasi baru / berubah, tinggal update dict ini.
LOCATION_INITIALS = {
    "HWR0001": "HSM",
    "HWR0002": "HPI",
    "HWR0003": "HSS",
    "HWR0004": "HKC",
    "HWR0005": "HMB",
    "HWR0006": "HPB",
    "HWR0007": "HCC",
    "HWR0008": "HBB",
    "HWR0009": "HBG",
    "HWR0010": "HCP",
    "HWR0011": "HPM",
    "HWR0012": "HJB",
    "HWR0013": "HCS",
    "HWR0014": "HPS",
    "HWR0015": "HSQ",
    "HWR0016": "HAY",
    "HWR0017": "HMD",
    "HWR0018": "HTS",
    "HWR0019": "HGS",
    "HWR0020": "HDS",
    "HWR0021": "HOM",
    "HWR0022": "HSC",
    "HWR0023": "HPV",
    "HWR0024": "HOS",
    "HWR0025": "HWR0025",
    "HWR0026": "HWRZ001",
    "HWR0027": "HMO",
    "HWR0028": "HMS",
    "HWR0029": "HPK",
    "HWR0030": "HLJ",
    "HWR0031": "HRJ",
    "HWR0032": "HTM",
    "HWR0033": "HJS",
    "HWR0034": "HMC",
    "HWR0035": "HKT",
    "HWR0036": "HKR",
    "HWR0037": "HSP",
    "HWR0038": "HTH",
    "HWR0039": "HIP",
    "HWR0040": "HNP",
    "HWR0041": "HBJ",
    "HWR0042": "HSK",
    "HWR0043": "HPW",
    "HWR0044": "HRX",
    "HWR0045": "HRS",
    "HWR0046": "HJC",
    "HWR0047": "HPO",
    "HWR0048": "HQC",
    "HWR0049": "HBM",
    "HWR0050": "HAP",
    "HWR0051": "HHH",
    "HWR0052": "HSB",
    "HWR0053": "HMT",
    "HWR0054": "HSO",
    "HWR0055": "HAM",
    "HWR0056": "HPN",
    "HWR0057": "HSR",
    "HWR0058": "HJT",
    "HWR0059": "HTR",
    "HWR0060": "HKD",
    "HWR0061": "HCM",
    "HWR0062": "HBP",
    "HWR0063": "HMBP",
    "HWR0064": "HTB",
    "HWR0065": "HCT",
    "HWR0066": "HPL",
    "HWR0067": "HMF",
    "HWR0068": "HMM",
    "HWR0069": "HGI",
    "HWRA001": "ADG",
    "HWRA002": "AML",
    "HWRA003": "ASM",
    "HWRA004": "WAR",
    "HWRA005": "ABH",
    "HWRA027": "ATB",
    "HWRB001": "HWRB001",
    "HWRB005": "HWRB005",
    "HWRB006": "HWBPB",
    "HWRB008": "HWBBB",
    "HWRB014": "HWBPS",
    "HWRB015": "HWBSQ",
    "HWRB023": "HWBPV",
    "HWRB024": "HWRB",
    "HWRB025": "HWRB HOM",
    "HWRB026": "HWRB026",
    "HWRB027": "HWRB027",
    "HWRB028": "ZBB",
    "HWRB029": "HWRB029",
    "HWRB030": "HWRB030",
    "HWRB031": "HWRB031",
    "HWRB032": "HWRB032",
    "HWRB033": "HWRB033",
    "HWRB034": "HWRB034",
    "HWRB035": "HARDWARE BAZAR SIDOARJO",
    "HWRB036": "HWRB036",
    "HWRB037": "HWRB037",
    "HWRB038": "HWRB038",
    "HWRB039": "HWRB039",
    "HWRB040": "HWRB040",
    "HWRB041": "HWRB041",
    "HWRB042": "HWRB042",
    "HWRB043": "HWRB043",
    "HWRB044": "BAZAR HMS",
    "HWRB045": "HWRB045",
    "HWRB046": "HWRB046",
    "HWRB047": "HWRB047",
    "HWRB048": "HWRB048",
    "HWRB049": "HJF",
    "HWRB050": "HPF",
    "HWRB051": "HLKC",
    "HWRB052": "HLCT",
    "HWRC001": "LMP",
    "HWRC002": "BJM",
    "HWRC003": "BJQ",
    "HWRC004": "LPA",
    "HWRC005": "CVO",
    "HWRC006": "LMK",
    "HWRC007": "HWRC007",
    "HWRC008": "HWRC008",
    "HWRC009": "BDM",
    "HWRC010": "HBL",
    "HWRC012": "HWRC012",
    "HWRC013": "HWRC013",
    "HWRC020": "HWBJQ",
    "HWRN001": "HWRN001",
    "HWRN002": "HWRN02",
    "HWRN003": "HWRO03",
    "HWRN004": "HWRN004",
    "HWRN005": "HWRN005",
    "HWRN006": "HWRN006",
    "HWRN007": "HWRN007",
    "HWRN008": "HWRN008",
    "HWRN009": "HWRN009",
    "HWRN010": "HWRN010",
    "HWRN011": "HWRN011",
    "HWRN012": "HWRN012",
    "HWRN013": "HWRN013",
    "HWRN014": "HWRWHJATIM",
    "QC3": "QC3",
    "RJ3": "RJ3",
    "WH03": "WH03",
}


# ============================================================
# LOGO HEADER (BERITA ACARA)
# ============================================================
# Logo HARDWARE untuk letterhead Berita Acara, disimpan
# sebagai base64 supaya file .py ini tetap satu file
# (tidak butuh file gambar terpisah saat deploy).
HARDWARE_LOGO_B64 = "iVBORw0KGgoAAAANSUhEUgAAAWYAAABSCAYAAABnun4ZAAAgAElEQVR4Xu2dCbQcRRWGGwENSkRDiIJRAiogi0oAUaMiKosLCsiegDuGBAQ1iBsCCkFkERDUGBEVCAaBBMK+Ci6guBD2PWxqNCQQjILK4vmaVJ+2cu+t6pnpmek39Z3jkZ43edOvp+qvW3er5Z599tlns0QikUj0BctBEuZEIpHoH5IwJxKJRJ+RhDmRSCT6jCTMiUQi0WckYU4kEok+IwlzIpFI9BlJmBN9CRmcyy23XF/eW6cZpL81EUfLwnzPPfdku+yyS3Fdhl938MEHZzvuuGPxWivst99+2a9+9atlJii///Wvf332k5/8pHitHfh93Ov999//f5/F68cff3y2xRZbFK+1w3/+85/8c/785z8v8zeF4F5WWGGF7EUvelH2/Oc/P1tjjTWyMWPGZOuuu262ySabZOuss07x3qr85S9/ye+L+2uFYcOGZSuttFK2yiqrZK961auy173uddkb3/jG/Dvi9Vb4zne+k/3oRz9a5jnxHFZdddXsvPPOy59Fu1x//fXZPvvsI37ON77xjewDH/hA8Vqr3HzzzdlHPvKR4roMn/OpT30qmzRpUvGaxBe+8IXs8ssvX2Z8Tp48OfvkJz9ZvNYJfvzjH2cnnHDCMs8klhe+8IX5mHjpS1+avfzlL8/WW2+97A1veEM+Hl784hcX76vKhz/84ey+++5r+b6q8Mwzz2Sf+MQncg2y4H1XXnlldskll2S33nprtmDBgnzMr7nmmrluvP/9789e8YpXFO+PpWVhvuWWW7KNNtqouPY5+eSTg4MtxJZbbpn94he/KK7L8EXfeOONxXU7IMhrrbVWcV3moIMOyr75zW8W1+3w73//Oxeuv//978VrnQDB5nnstNNO2Uc/+tF8MlTB+vvbgYUDYRs/fnz25je/uXg9hpkzZ2a77bZbce3zhz/8IRs7dmxx3Srf/va3s8997nPFdZkDDjgg/3m7MBf23Xff4tpnxowZ2e67715c+/z3v//NnyULqM+2226bXXzxxcV1JzjiiCOyr371q8V1p1h99dWz97znPbngtWLsIHYPPvhgcV03jItjjz22uPbBODjkkEOyuXPnFq/5rLzyytnEiRNzQ7XKotSyMN92223ZBhtsUFz7/OAHP8gtgXZ43/vepw46Jvp1111XXLfD2Wefne28887FdZl3v/vd2RVXXFFctwPCjPWAENbFaqutln3pS1/KPvvZzxavhXjggQdyy5v7qwss8q9//evmmCnDM+JZaff0ve99Lx/w7bL33ntn06dPL67LvPe9780uuuii4rpV9tprr+y0004rrsuwqN55553Z2muvXbzmgyWGtYl15oM1dvfdd7e8M5H41re+lRskdbLddttlxxxzTKWdHjuxO+64o7iumy9+8YvZkUceWVyXYQdz9NFH5/+9/vrr5wvrW9/61mzUqFHZv/71r+yuu+7KzjnnnGz27Nn5ezBiZ82alb361a/Or0MkYc6ybP/9989OPPFE90z+jxEjRmT33ntv9pKXvKR4rVW6IcyObbbZJjvzzDPz7WSIbggzsMVlMuI6CIGtwGBGlCT23HPP7Kc//Wlx3SrveMc7sl/+8pfFdZnXvva1uQGCeLYKYspuhh2mBAvVTTfdlD3vec8rXvM59dRTs49//OPFtc/vfve7bLPNNiuu26Ubwgy4vjDgNJeoT7eFGQNn6tSpxec7EOyjjjoq/2/+H5HW+O1vf5sbqLizEOXf/OY3uXiHSMKcZdlb3vKW3Neoce2112Zvf/vbi+tW6aYww5ve9Kbc9xUS524Js0Mb8D64ZbQ4AhY+or388ssXr1Xln//8Z26xSS4CwJePhdaOm4dny2do/vsJEyao1rQDUUacNb773e9GLXaxdEuYHaeffnru7grRD8LMDp6dPH5u7nuPPfbIlixZki1evDg3JvgfPyP+4ebdY489lrtwcL996EMfKqxoi4EX5vnz52evec1r8kmqEVoVY+m2MAPBhwsuuGDpHch0W5iBbeCUKVOW3oEMLgZcDRIIMqLJd9cqWLG4CCxPHq4MXBqtwiTcYYcdimsfdmpWgIl74x41ixsQNUSiU3RbmNktYBiFrP5eCzO7HwLtxLbwwRMcBr6/U045JZ/fLjDJos6u9Yc//GEerJ43b17+PSLiBAvf9a535e/TGHhhvvDCC4ORd/xh559/fnHdKr0QZmDQWFvhXggzA5iMG/xyGn/605/yiaAJ589+9rNs1113La6rgs8vlDmE6+Xzn/98cV0VdgdW8JitLjsbDSY0vnbN4gYWp9tvv70tl0uZbgszbLzxxtkNN9xg7oA0YWYsaWOkHYjTHHfcccU1bohx48blcRzcm8OHD89fdzs7FmFcn9zLueeem2e2EBw89NBD8/cRTOa1mF3SwAsz0dLDDz+8eCASo0ePzgMspAC1Q0iYcZcgoNog4/WFCxfmg4JB/Mc//rH4mQUBIgY0EWKJkDCT6kXGhxR8YlL84x//yH8HrgX8tbGRcyyI3//+99mKK65YvFbmiSeeyP28pBdK+BOnKlg8X/va14prCf520sdaheDxVVddVVyXwddI2qmb4BJWYNqBmPHs+Q47gSXMiD9uKMRJGqdYv4wH3EP4VRGzRYsWFT+3ID3yYx/7WHHtYwkz2SlkgUljtFVY8AjsOQgEfvnLX87vkXt1OGEuPw/mEnpRFmFiAZtvvnnua+bvsBbSgRdma+KUQQRZ1dshJMxsibQgpMSvf/3r3M0yZ86c4jUNxEXLpQ0J8/e///3s05/+dHFtgUsIS5SUK2kS+YSs+TozcxA8hM9i0003zRfBVsDvyOR+5JFHitfKsJ1lW2thBabLEAglINoJLGFGbEj3tBaTMrgKES12DfhaLdgdsVBrIJLsDCQOPPDA/L7rhCAergl/F+VeZ6HHbfH000/nO2zyzi+99NJs6623zt/Hs3jlK1+ZCzJzzgoCDrQwM1BYvWJWdAoerFzUGELCzBdMlLoqiDORYguCDwwUiZAwY5VWSb8DrCYENyR8ZF7gstC2sKTYsR2UQBywOK0BbkERjJWDCkw0dkuhAKoEgm65KWKCoPhdLbFytDp2JCxhfsELXpCngpGPXwUEFZcguz0LxgLfi4RmMUNVo6YVnABTdMaC6SBtc9q0aXmQmEUL44TdIAHbcq49P0OYARcVRWIaAy3MbLtJl4qBYgfSz9qhLmEGkuGtYgjcGEwKScTqEGbH9ttvnyfiW/A9vO1tbyuuy5BDvtVWWxXXPvycXU9VqNDCTYJVG6LVYhayJajM02BnwfPRiAlMO9jNxbq2QtQhzEC+Nlax9fewEONelOi1MDu3J9/pSSedVNyXy5rh73rqqafymBV/K+O+XFjlaj8QZOai5RodaGFme4XVEkMnclrrFGaS2rlHLfULSJ0jUuxTpzCzK8EC4jM0LF/x3/72t3xXo01mXCb4/aqCC8QKPJY544wz8rSoquA60nKtidozeano03CpWTFQYIJgEg9pl7qEGdj9IL4abPvZ/kv0WpiZP2ToMM+w3F3uOT5nXIXsvNlZMQdxu/BzntXIkSPz91FFSCZSTKbUQAszDyi2souHjDC3E2CpU5iBQJZL4ZEgeCG5POoUZsDHSIBEA2sUq1SDgAmBE4lWq/Pwbcf2mCBV0hUUxEIgyCqQYatL2pVLr5JgwdEqzyTICiBPtl3qFGZEi10AgV0Jyq5ZsPgcn14LczkYXQ5UulgFuzAnwi7jx8URmFvksxMYj8kmGlhh5iFjif31r38tHkaI2ER4jbqFma0sW0UNxFEqVKhbmPm9PGstuwJrDz+u1uwF/2F561iG/gv826oNjVzqUgytiD/fMd+19kz9yL4E/SQoboqlUwGwOoUZrB44/H6sUWknYQX/uuFjBueewjIm+M5igTuNe2ahL7sn6IFC4JfiH+Y2BgoLMkZIaOc9sMLMwyHiXgUyE8hQaJW6hZmAG9aI1iSJCSFloNQtzMCzs/42y1fMALcWxFYyZrBkrr766uLagkUFQdDS+iRCaW6har1HH300758RymQog2sGsWiXuoU5tCiSYkc1rk8/CDPgQ6b+AV8xxhrzysL1SuHZ8f1YxpNjYIWZLIvPfOYzxYOIgdWOqLHV18CibmEG+jLQe0Fiww03zH/mb5+7IcyhbnFWUyK2tgRNSEOSqJLOB3wPbEkfeuih4jULJhT3wDY7llBhCZkW1gRl0QhVh/nQe4KgEpkk7VC3MIe67ZFqRgaHT69dGQ4MINwUrrkZux/ylYml8B1QDMTYIkbAAkQGBn1iWKxjq0gHVpgJ5lTNsmBQsm12KS9V6YYwk8GgdcNjMjGp+DvKdEOYWRAYuFJRApQrpHxoe8mCRi9eiaolyXyHWF9E0GMp56PGYFnktGVFQJmsGqTRfeUrXymuY+Ez3/nOdxbXrVC3MJOtYGWjkJJGa1CffhFmoJCF8UprgSeffDJ/DXcaWU+4ScmocTAWCPxpaYASAynMTHS+5FBOpQSlllbvA4tuCDPNxLlHCQQBy8/vC9sNYcbXhpWqbc1DbiKrGATRJsgWu5OhIOeDH/xgcR1DFXcOFhXuD4JBEiyel112WXEtwf3FFA75kM7ViqCXqVuYQ2mq5ARLPVL6SZgdaAiuNhZu3F2kX5JxQ3YMlbx0zpMyoUIMpDDHNK/RCDXPtuiGMFM6TR9YCQIWCDPltGW6Icz87UwstnUS+OG0TnJgNbRnIjAprJ7GZaqkSTqqlGbT/8I6GMDaHQCpj8QKqgSmHa0EKn3qFmbXc0JDc2v1ozCXYUEmrZPAHu4k32VYhYEU5phUKa0xCoUQWv/eEL0WZnpKM6l6Icy4DXAf4EaQCAkfDY+s1qtY0+wWYqCpOSlLVaCCD8GNIRS/CHWss4SdcUkQUmtqxFYaK07rixJD3cJMAEwrKAItMIoxRf8NiW5lZXSLgRRmgmz4sSTINWTQaKlMuAEoA/bFLYZuCLO15eeesZj98uJuWMwh91FImHGBYBGTrSBBiSylsiFYbMngCJVi+2AB8b3HHJiAz5vtrQTpVAibFacgYETmggSfz+7CKj+meEYT9hjqFuaQxawFc5mXWtZJEualNPVoKZz2OOG1lRdrii5VlnWmVdCF6IYwWyXQiAEWay+Cf+0KM1h5vaRXMeFD4PfF/8u2syoxpdmMLzJItJ4OZM2Q2WNtc/FL/vznPy+uy2A1nnXWWfn/a1Yzwm5Z7CHqFuZQxonWZS4JcwRNFWYi+wiENqiJhuPLxKrUqpOsen6LbggzlogmUFSiYSn6otAEixkQC0RDInYnQ0N2KUc2htChqRAaX6HCEhYwcso1XzwVY7hhrJxehJ30xFapW5hDeemUK1OV65OEOYKmCjOWiHXGmEuLovhEKxOOiapL1C3M/H5X9imhdZhrijCTbWL5ka0iFQeiKKVixRBTWRe6x1CbU+YVC6jWV9ideELeLD08JAgc8nuqFMSUqVuYXV9jDRZPyvB9kjBH0FRhZotHcEaCPES22i972ctydwbRYQmsMt4X25PWUbcwM2HYRmv5uQQ8pROhmyLMJO2z8Li8UR/6hHDkj0Vsf2OJmIyH8unJPuxU6I+BG0IjtHA4/zF+WClABnwOrrrYE8l96hZmy1XDHOT3Sy0xkzBH0FRhthri0PvW/YziE6ujmLaqW9QtzKHTlP0+so6mCDOBO3y0WnwgpmtX7MEIEuRh4z7QekeDVVgScxKO61QmUc64CLUUYAEOZR5p1CnMpAISxKVroARuHO2Q3STMETRRmBkMTC4t8FOO7DP48ONpZcD+KQYx1C3M1sAF7bTvpggzWMJFEyS+N62iDlHA4taaKRHYoykSfRAkEFSEWWqwAxQXEFjk+C+JkMWN+wJrWutIV3ZFhXKdQ75sizqFGfcLbhgNq++5Nb5TVsZSmijMof62DAjXzwF3AFtBBqEE2Q+09qtCncJsbW0BweFvkfJbmyTMPB8plcphNTTC0sbilvLTgWoz+olYGQ1WaXYosBgKGhM4ZHwQAJSgKKV8motrpiOh9UWJoS5hJuWR/iBaaT3w/TIPJJIwR9BEYaZUVTvKh0AJKU7l6jHOUNN6MGA18X4Gaix1CTNNX8hf1jIBQGv5CU0S5lBrU604AUId36guxJ2FAGhYOyV2W1bZtiXqQBqc1afXT9PkkADNp854xm1SpfGSow5hZjeBb9kKmrMj4Z61Zv9JmCMICTPbTe3wz1joMKX5DFs5iNPy/5HiRKl2ud+C1QUrJpDjExJmAo58ZiwPP/xwnrOKWISwMhaaJMyhLbxV2h06SADhxKLGJaK5sCwXgVVRSM9p0vmkoJbDaocppQOG8oEJsFGiX5WQMOMKiu1g5w4m5YAGbffpoMk/zf41LGFup1VCP1Jb5R/NVJgk2gB3WFstLFat/LmqMFMxhv9Pqxzjs/xjgLDOCLBoj0frgqUREmZaCbpjd8qfyTPiOWJ14L/EUmfbTLPxmDPrWDysdqVNEmawTp6RFliH1eCJ9/NciUFgsWl+6HKAuAzfD2luWm5xzGnblH1r75HGO42hWKS0MYBLRhN6C0uYscRprsQz8ucFz5D4zeOPP567K3ClEGjVim18eK/V29gSZlyQ7B78e7LgvSx0NPfqN2oRZoSEqCrNPEK/2v3cF2iuETLt30sD1SLU0UoqA6XAhIGvnaNnuQckQsJcF1rCvqNpwnzYYYepTYAQDsrOObG4TKgiDysZi5TttHXCBpYimRH03S3DvyVYrPmHQ7shdj8EJrWiJk1kCebSR0SCrCEW8KpYwuxgfmpzsxVw0eCqsdCEuZ17wedvnUHYK2oR5m5QVZhDg007Nt1yp1ipPRK9EGZpJ+DTNGEmMyHkq/V9yVjACB+uEIny6S5WrjMiwFjB5VEm5B/G/SGVGTsI4hHM09DcEoxpxraEdTK6RWiudBrqAdidYgRZaMLcDkmYO0xVYbb62xLIwP8lBfKshuUIMhZYaEA5ui3MpFfh38O/adE0YeboLIK02snZkr+RAzF5Hhr4dwn+Qaj7oFSazWe6f++DmFMKj6tDwzrxRApMO/h+rQNYQwFHiW4KM6LM87QWJUcdwkzcgR1YvzEQFnP5dFsJK+iAFaUFzUCapBrdFmYrQ6FM04QZLH+sdPad1bENyqXSVttNIJDln17NqSHXXHNNcV2G5lEs/FZhidWgidQ3As3SzixUDYnoID5V6KYws1hpR6H5JGGOoEmuDM5XI2ijwSRjskkQbEPUtaAhWRtaibdPt4WZBvIEyayFBZoozFbJvHT2HamIWqtXKFdy0oGOXRBBLAly4cv5w+Tn8n6tsAR3GJatxqJFi/LAtHa6i1ZKD3ghrTamWn8Ui24KM7jGTCGSMEfQJGEONS4PRYMta0aL0kt0W5iBiDN+8BEjRiy9i2VpojBz6jBZPxr+2XdY0ZrrSxJy4g2a2CHCjH/XJCjUXxh3mHViSijtLZT9Q2GMJtwjR47MMySq9HXptjBDzO7OCnS2CpkcVgplrxgIV4aVX0p7Tyal3zy+jBVgsZqu+PRCmKHsP5VoojDjc2WLr6VjlndBS5YsycVU689AwQq7qjLWYb1+kUWosCS08FuFIoAbww82lgktUrh1WJhi6YUwYzjwTK38aEuYcfOwQ6ySx8B7GSNahk8vqU2YaetHbmHsry6/z6XOkb6mpfvEBv9CRxphDWupUY5QgAX/tPVzR0iYKfMmdYfULvcMfAh4UVxBri73rRXMlKE4gUFP1zyJJgoz3yvPkkVVgmCSC/aGTuiWilKsoC9QwUb7V2ArTlaGRExmBK4R2gVIkC/Md2cFcJmLVqtQhJaWpbFYwswugVYELErl5+n+m3HLd0OONZY6uwmelbYoliH4qX0uWK4MitmwfLXvWILnxU6CRUGbb72iNmEmRYtUrXawMilihZltPNaGZlmxgGCxWNDfmAALwirBNlUr9S4TEuZQrqsEA3/KlCnBZ3HSSSdlkydPLq7LNFGYwTpGi7xkFmMELdSDmzadPMMy+Oat3G93ajbjioUf8ZQI5RKz0JJtQaaJBCeya0UxjtDpMBQuaQf0SljCXLXyD/ChE3wN5QuzuJCKKAU5wRJmvgu+k6FCbcJsNSOJxbIkYoU5lPqE1YnVHILmNNoEo3BFi8iXCQlzq70yyDrZdtttVT84WAn8TRVmStEtS9A1NMJNYC2+Url6qKGQC8hRWIIoaj2wQ8HhUAZIyL/s4D1aqTj9Mii6kdJBJULCXHbjVIFdCUVZFvj1tTYHljCn7nJLaYowM4EQZwm2L+S7+tsyH0pNeZ92ZBMHZDKRLT811CXMgLjic8WfKsF2DRGR7rGpwhyq5nTGAW4mLSuCFqEIDRZ2GRpCYQlrVqgzDEKNkUI7RwpZpB7ZwPjEnWL5lx1YxFocBbQCKom6hBlwOVgFTyxiWn+aJMwRNEGYEVsGtdZYnYFvCXIVsJgtkYA6hRmsbniAdUb+r09ThTmUpkZeMosybigtxkD1JmNZ6q3BLoQCDQl6LOB+sApD2JITC+A717BO86g6Pq33x2Q9OOoUZhYIgq3afUr+fkcS5giaIMxYsdbBmJ2ErbJ1jhnULcyh8+ZoUi6dytJUYQarrwWNgzgxHGHWqgQtH651TBTgHiDjRXO3kZtM9gg9YyQYl+TIa2c0dpKY0nxHncKMT545qS2UZF5oLjkrKyO5MpbSBGEO9S/oJKHTKaBuYaazGQEULdCp9RJusjDz92hBHwJUWLM8Vw2rJDeUhkaBCy4uXEQS9LbQrGEIZYt0kip9XeoUZuC5aMFIFlG0RbrPZDFH0ARhtprRdBpOCGGCascaQd3CTPQbC4z/l+D0Cylns8nCbO0S2NojIPx9GviItX8fOlePZ83v1nZkUrZHGb5rv6NhXWC1M2e55xB1C7N1IDL1AOxEpJN2kjBH0ARhtjIp6kDz4TrqFmb6JbBN1H6/llLUZGHmb8XK0rInLPArE38gyCcR6uEdIpTxY51fWAflo9Ms6hZma5dDpSLCLFWqJmGOoN+FmYR2AkNalkIdWBFl6IYwI/yahYh1Rt9pnyYLM24bArzaAaYWNBdCBKziDbIvWHCrQvYLflQt35fiBrJotMb6dUAeO/nsIeoWZgwEd+ixD98Fz83PkoEkzBH0uzCHCgTYMpHfqVVLaRBl1wJJ1gm/0GthJipPdN6nycIMoRQsjZgGP+TdalkCFoiIdvoOkIbH7kaz9LEcMSyq+J/ZAbDQaK4scrrJ7Q7RS2HGFYgwSy0OkjBH0O/CTANsjrfSwDdJRL4qVkoaFVxYQNTsSwy6MOMy0AJliGur2/pp06ZlEydOLK5jiYnkh/pgaGhuIwcLuJQh40AcreIZDavvhmWNlknC3HuGbOUffQyo6JJggJJK18pZX7gCtHxQgk1Y1JrPcpCFmbPg8AXPnz+/eB5lrPzVEKGTszU4FixUiRZqsK8xc+ZMswycMSS5lRwUMxEjqUqofzjtSqmotUjC3HuGpDCHgjYxB2NqhHo7W5N9kIU51NA9VLpsETqbUYN2rdZ3Cdw3qWbaWXwSZECwc+KeNBiDZH1I0GyKnYWUmRCC5kHs3FgIJWJ6wyRh7j1DUpiJhlttFkmj04IPITgzjglHhzcJV20mMcjCHCqfxu1kdXMLYVXpSVBGj/hpwTlHqCBCgvFBYYmUiwuh8weJjWjnTMZg+WJjuikmYe49Q1KYKSqwGpNbuasxWM2ViLRTOIBbw2eQhZl+0JyLp4Hffvz48cV1VcjP1gpFJEhrjM22sHptSIRO5KAakRavGjFVpBZW1gMLEoFHKR3NkYS59wxJYbZOtqafLOk+Y8aMKV6rCgIgFWoA6URExsn48BlkYbZ8/hDjVrDAWsZqjgV3E26nGELd6XwQRa0xEZCKhvhp+KevVAX/tpWvHPr9SZh7z5AT5pC/kZaCNFKRmtbEEhIBrCWpFHxQhXnevHl5frVWIYc7ASuOI55aBdcS37vmHvChP7B13FiZqqX9ocAdLh0tlS7mRJ0QGB4EoLXS/JBFnoS59ww5YcbycodqSmhFFlVABCht1fKZmfBMfJ9BFeZQhRvn3ZH90C5Y3P4RURqhqrwyVXpasMjgu8ZlIEFwjsC0dtBrJ55FyC8e8mEnYe49Q06YQwevzpgxIz8DsF3Gjh2bW94S2qkVgyjMpGdxzJPFUUcdlXdyaxcq26QCGh+riEEidGZgmVBwLZR+RwDUyr+PBX89Y12ifLqLRBLm3jPkhNk6f40oOdFyK40pFgoaKGyQ4IBWtqP+GXuDJsy0b0SUtdQtILWMcmqyFNol1A3OwTafHhlV3FlWpkMZikIs/7FVAAKcMsNpM+1C6TUFNBqWTz8Jc+8ZUsLM8T5McHyaEvwMEdD641aBLALrZAr80FtvvXVxDU0VZjIq6DscC2XuuItoVqPlLTs415EshU7Ad0sMIVRmX/UMPLAW4jK0+aStpQaiy+GkEhyaixuEJvztwiLCYqJh9XVpojBr7sOmMqSEOdSPuJ2yXx8+i9Q4TQQoCfcPn2yqMFO2bFlfgL+dakoqz6jg09w8PhzPZJ15VwWCiwQZtYXZobU/tQhZoEDGDxk5a621VvFaGVwi+Je1g1fHjRunNoKvCj5sPguftgTuPM3V0URhxnDAgBgqDClhJv2JAg8NDqskENUJQiIgBXGaKsyk/tFNTFqEcAfgqliwYEGeCRMTIHPQK4JTVToJ1vCsWbOKa4lW+qTglgkFCwm4YbVLOexA3MHK1mCHwWEGnYIxSGqchNU4v1+FGd+8P6cctFfARVll/Fkw1jHkutUv22dICXPo4FWrj0UrWCJAdzD8zGxPHU0V5jog8HbjjTd2ZNtehub0ViAx5hw+CTJxcIVZbWQnTJiQ+7k1OO3EapzfyoJhQSBx6tSpxXUZngPCjED79KswV63ubBcKovjOesGQEWZWOPryMvY/b/IAAAMaSURBVOkk2F6yzWS72SlCIoCVxTlljiTMzzF8+PDcz9opF0YZXAHlZ+7DgkBGgnXSjASWGG4yxEzD8tsCp2lTdSoR2/mtClQrUrWogctJCpb2qzBb/vk6oHpYW9jqZsgIM/5NrCCtv22oTLYVQttbBni5dWMS5iwbNWpUbhniT60DehHjW+UEbYlQOpsFZfzawa1g+cv57nF1aK4vMiTIlOgk9OQg315rwKSdmJOE+TmSMCtYPSmYAEwER6gMleANea6dhO51DPyFCxcWr5Xxo/+DLsxsRU8++eS8+1mdWKlt2mEBMZBfTFBXApcMGRVl11WZuXPn5o3qNR9oHVkFfBa7SFIDJfgZQVrfJ56E+TmSMCtUEWarcQsDjxaLTIxOY5XXjh49Oi+PdYn8gyrMCCW9I6w0sk5i9aJgYZg0aVJxXYXZs2erPmB6IFu9QELN/Mm9x9XRaawTWLS+LkmYn6ORwswqTM6oRkyKVQjL8mFbyCkqQHocQQyCbRKUxxK8GTZsWPFap7AOlgR8nm7bTk4vp2lr22wrhSkGtqz8/sWLFxevldGa0bO97qQVS+cy3Ep8fxSYWD7fOmCXoi0CoQY+FhQnMe4kQoUl5GvPmTOnuPZh7HbyO3CETuLmSC4/H99q0gWMl3aagO29997Z9OnTi2sfemBj1PhYelAH7bQHbpeWfcykRyEi/jYICMSR2kKebztQeMAg8Cu0uF2q6vAb8/kIHoLD//v3w71gEWiWTrswWanW8u8R+GyKTJjM3BcFMETtyTGV7pPgknX6RAj86/x+0tek389CSgqVD0LOv+M9VeHfsCMgoMcCSHCNSUVWSq/AtUQaHuPEPQf332ROWC0vLRhfCBkLYPn58gwITGmizWczV8gp9r8Xfob7g0VTSl1rF06Mwc3nfy5w36Tv+b1l8HXTiMkf09wrxVncK993qxCb4dQZ6fcTnOf3S4cEoAecii79LZ2GZ4Pfv65YSIiWhNkN8kQc6Xn1F+n7SMTSq7HSkjAnEolEoj6SMCcSiUSfkYQ5kUgk+oxcmPvsnhKJRGLg+R/Xs38kQy0DbwAAAABJRU5ErkJggg=="


def format_location_label(loc_code):
    """
    Tambahkan kode Initial dari MASTER_LOKASI di belakang
    kode lokasi, mis. "HWR0019" -> "HWR0019 ( HGS )".
    Kalau kodenya tidak ada di master, atau Initial-nya
    sama persis dengan kodenya sendiri (tidak menambah
    informasi), kembalikan kode aslinya saja.
    """
    if not loc_code:
        return loc_code

    initial = LOCATION_INITIALS.get(
        str(loc_code).strip()
    )

    if not initial or initial == loc_code:
        return loc_code

    return f"{loc_code} ( {initial} )"


# ============================================================
# MASTER LOKASI - NAMA LENGKAP
# ============================================================
# Sumber: MASTER_LOKASI.xlsx, kolom "Location Name".
# Dipakai khusus untuk Berita Acara, format "NAMA LOKASI
# ( KODE )", mis. "HOLY WOMAN PLAZA FESTIVAL ( HWRB050 )".
# Kalau ada lokasi baru / nama berubah, tinggal update dict
# ini (samakan dengan LOCATION_INITIALS di atas).
LOCATION_NAMES = {
    "HWR0001": "HARDWARE PLAZA SEMANGGI",
    "HWR0002": "HARDWARE PONDOK INDAH",
    "HWR0003": "HARDWARE SUMMARECON SERPONG",
    "HWR0004": "HARDWARE KOTA KASABLANKA",
    "HWR0005": "HARDWARE METROPOLITAN BEKASI",
    "HWR0006": "HARDWARE PLAZA BINTARO",
    "HWR0007": "HARDWARE CIBINONG CITY",
    "HWR0008": "HARDWARE BOTANI BOGOR",
    "HWR0009": "HARDWARE ISTANA PLAZA BANDUNG",
    "HWR0010": "HARDWARE CIPUTRA PEKANBARU",
    "HWR0011": "HARDWARE PALEMBANG SQUARE",
    "HWR0012": "HARDWARE WTC BATANGHARI JAMBI",
    "HWR0013": "HARDWARE CIPUTRA SEMARANG",
    "HWR0014": "HARDWARE PARAGON SEMARANG",
    "HWR0015": "HARDWARE SOLO SQUARE",
    "HWR0016": "HARDWARE AMBARRUKMO YOGYAKARTA",
    "HWR0017": "HARDWARE SUNCITY MADIUN",
    "HWR0018": "HARDWARE TUNJUNGAN SURABAYA",
    "HWR0019": "HARDWARE GALAXY SURABAYA",
    "HWR0020": "HARDWARE DELTA SURABAYA",
    "HWR0021": "HARDWARE MOG MALANG",
    "HWR0022": "HARDWARE SUNCITY SIDOARJO",
    "HWR0023": "HARDWARE PEJATEN VILLAGE",
    "HWR0024": "HARDWARE ONLINE SHOP",
    "HWR0025": "HARDWARE ONLINE",
    "HWR0026": "HARDWARE ZALORA",
    "HWR0027": "HARDWARE YOGYA MALIOBORO",
    "HWR0028": "HARDWARE RATU INDAH MAKASAR",
    "HWR0029": "HARDWARE PANAKUKANG MAKASAR",
    "HWR0030": "HARDWARE LIPPO JEMBER",
    "HWR0031": "HARDWARE HARTONO YOGYA",
    "HWR0032": "HARDWARE TOWNSQUARE MALANG",
    "HWR0033": "HARDWARE JAVA SEMARANG",
    "HWR0034": "HARDWARE MARGO CITY",
    "HWR0035": "HARDWARE KEDIRI TOWNSQUARE",
    "HWR0036": "HARDWARE LIPPO KARAWACI",
    "HWR0037": "HARDWARE SUN PLAZA MEDAN",
    "HWR0038": "HARDWARE MALL THAMRIN MEDAN",
    "HWR0039": "HARDWARE PALEMBANG ICON",
    "HWR0040": "HARDWARE NIPAH MAKASSAR",
    "HWR0041": "HARDWARE BG JUNCTION",
    "HWR0042": "HARDWARE SKAMALL PEKANBARU",
    "HWR0043": "HARDWARE PURWOKERTO",
    "HWR0044": "HARDWARE ROXY JEMBER",
    "HWR0045": "HARDWARE ROYAL SURABAYA",
    "HWR0046": "HARDWARE JOGJA CITY MALL",
    "HWR0047": "HARDWARE PAKUWON SURABAYA",
    "HWR0048": "HARDWARE QUEEN CITY MALL",
    "HWR0049": "HARDWARE BATAM MEISTERSTADT",
    "HWR0050": "HARDWARE ATRIUM PLAZA",
    "HWR0051": "HARDWARE M BLOK",
    "HWR0052": "HARDWARE PAKUWON SOLO BARU",
    "HWR0053": "HARDWARE MANADO TOWNSQUARE",
    "HWR0054": "HARDWARE SOLO PARAGON",
    "HWR0055": "HARDWARE ARTOS MAGELANG",
    "HWR0056": "HARDWARE PANAKKUKANG NEW",
    "HWR0057": "HARDWARE SUNRISE MOJOKERTO",
    "HWR0058": "HARDWARE JAMBI TOWNSQUARE",
    "HWR0059": "HARDWARE TRANS MAKASSAR",
    "HWR0060": "HARDWARE KENDARI",
    "HWR0061": "HARDWARE CIBINONG CITY MALL",
    "HWR0062": "HARDWARE BLOK M PLAZA",
    "HWR0063": "HARDWARE PAKUWON BEKASI",
    "HWR0064": "HARDWARE AEON TANJUNG BARAT",
    "HWR0065": "HARDWARE KUNINGAN CITY",
    "HWR0066": "HARDWARE PALU GRAND MALL",
    "HWR0067": "HARDWARE MEDAN FAIR",
    "HWR0068": "HARDWARE MEGA MALL MANADO",
    "HWR0069": "HARDWARE GRESIK ICON",
    "HWRA001": "ARL DAGO BANDUNG",
    "HWRA002": "ARL MALANG",
    "HWRA003": "ARL SEMARANG",
    "HWRA004": "HARDWARE WEB ARL",
    "HWRA005": "ARLBYHARDWARE",
    "HWRA027": "ARL TAMAN SARI BANDUNG",
    "HWRB001": "HARDWARE BAZAAR PLAZA SEMANGGI",
    "HWRB005": "HARDWARE BAZAAR METROPOLITAN MALL",
    "HWRB006": "HARDWARE BAZAR HPB",
    "HWRB008": "HARDWARE BAZAR HBB",
    "HWRB014": "HARDWARE BAZAR HPS",
    "HWRB015": "HARDWARE BAZAR HSQ",
    "HWRB023": "HARDWARE BAZAR HPV",
    "HWRB024": "BAZAR LMK 2",
    "HWRB025": "BAZAR HARDWARE MALANG",
    "HWRB026": "BAZAR HARDWARE PEKANBARU",
    "HWRB027": "BAZAR SOLO",
    "HWRB028": "BAZAR BOTANI BOGOR",
    "HWRB029": "BAZAR DELTA SURABAYA",
    "HWRB030": "BAZAR PEJATEN VILLAGE",
    "HWRB031": "HARDWARE BAZAR SERPONG",
    "HWRB032": "BAZAR MATOZ",
    "HWRB033": "BAZAR MALL KARTINI LAMPUNG",
    "HWRB034": "BAZAR CIPUTRA SEMARANG",
    "HWRB035": "HARDWARE BAZAR SIDOARJO",
    "HWRB036": "BAZAR ATB",
    "HWRB037": "BAZAR MAG",
    "HWRB038": "BAZAR MALIOBORO",
    "HWRB039": "BAZAR PARAGON",
    "HWRB040": "BAZAR HPM PALEMBANG",
    "HWRB041": "BAZAR HKC KOKAS",
    "HWRB042": "BAZAR GRAND INDONESIA",
    "HWRB043": "Bazar TP Surabaya",
    "HWRB044": "BZR MAKASSAR",
    "HWRB045": "BAZAR RATU INDAH MAKASSAR",
    "HWRB046": "BAZAR BIP",
    "HWRB047": "BAZAR PANAKUKKANG",
    "HWRB048": "BAZAR ROYAL SURABAYA",
    "HWRB049": "HARDWARE JAKARTA FAIR",
    "HWRB050": "HOLY WOMAN PLAZA FESTIVAL",
    "HWRB051": "HOLY WOMAN KUNINGAN CITY",
    "HWRB052": "HOLY WOMAN CILANDAK TOWNSQUARE",
    "HWRC001": "HARDWARE LAMPUNG CITY MALL",
    "HWRC002": "HARDWARE BANJARBARU",
    "HWRC003": "HARDWARE BANJAR QMALL",
    "HWRC004": "ARL LAMPUNG",
    "HWRC005": "HARDWARE CARVIENO",
    "HWRC006": "HARDWARE LAMPUNG KEDATON",
    "HWRC007": "HARDWARE BZR HO",
    "HWRC008": "HWR BAZAR HO BARANG RIJEK",
    "HWRC009": "HARDWARE BANJARMASIN DUTA MALL",
    "HWRC010": "HARDWARE BLOK M",
    "HWRC012": "HWR BZR BINTARO 2",
    "HWRC013": "HWR BZR BINTARO REJECT",
    "HWRC020": "HARDWARE BAZAR BJQ",
    "HWRN001": "BAZAR HWR LMP",
    "HWRN002": "HARDWARE GUDANG SAMPLE",
    "HWRN003": "HO MARKETING",
    "HWRN004": "BAZAR ARL LAMPUNG",
    "HWRN005": "BAZAR MALL KARTINI LAMPUNG",
    "HWRN006": "BAZAR BANJARBARU STREET",
    "HWRN007": "BAZAR BANJARBARU QMALL",
    "HWRN008": "BAZAR ARL BANDUNG",
    "HWRN009": "BAZAR QMALL",
    "HWRN010": "BAZAR HWR ALL BRAND EX PIM",
    "HWRN011": "BAZAR HWR LMK KEDATON",
    "HWRN012": "BAZAAR ISTANA PLAZA BANDUNG",
    "HWRN013": "BAZAR ARL BRAGA SEMARANG",
    "HWRN014": "WAREHOUSE AREA JATM",
    "QC3": "Gudang Barang QC HWR",
    "RJ3": "GUDANG REJECT HARDWARE",
    "WH03": "WAREHOUSE KELAPA GADING HWR",
}


def format_location_full(loc_code):
    """
    Format lokasi untuk Berita Acara: "NAMA LOKASI ( KODE )".
    Kalau nama lengkapnya belum ada di LOCATION_NAMES,
    fallback ke format_location_label (kode + initial).
    """
    if not loc_code:
        return loc_code

    code = str(loc_code).strip()
    name = LOCATION_NAMES.get(code)

    if name:
        return f"{name} ( {code} )"

    return format_location_label(loc_code)

REQUIRED_COLUMNS = [
    "OpnameNo",
    "Date",
    "Loc",
    "Golongan",
    "Artikel",
    "Barcode",
    "Qty POS",
    "Qty Fisik",
    "Price",
]

# Kolom yang tidak wajib ada di file input.
# Jika tidak ditemukan, akan otomatis diisi 0.
OPTIONAL_COLUMNS = [
    "HM",
]

# Alias nama kolom mentah -> nama kolom standar.
# Dipakai kalau nama kolom standar (mis. "HM") tidak
# ditemukan persis, tapi ada kolom lain yang maknanya sama.
COLUMN_ALIASES = {
    "HM": ["Cost"],
}

DETAIL_COLUMNS = [
    "Barcode",
    "Artikel",
    "Golongan",
    "Price",
    "Qty Fisik",
    "Qty POS",
    "Selisih",
    "Total",
    "HM",
    "Total HM",
    "Remark",
]


# ============================================================
# HEADER / COLUMN HELPERS
# ============================================================

def normalize_header(value):
    """Membersihkan nama header."""
    if pd.isna(value):
        return ""

    value = str(value).strip()
    value = re.sub(r"\s+", " ", value)

    return value


def canonical(value):
    """
    Menyamakan format nama kolom.

    Contoh:
        Qty POS
        QtyPOS
        qty_pos
        Qty-POS

    menjadi:
        qtypos
    """
    value = normalize_header(value).lower()

    return re.sub(
        r"[^a-z0-9]",
        "",
        value,
    )


def find_header_row(raw_df):
    """
    Deteksi otomatis baris header.
    Mencari hingga 50 baris pertama.
    """
    # Kolom opsional (mis. HM) dan alias-nya (mis. Cost)
    # tetap ikut dihitung skornya supaya baris header lebih
    # akurat terdeteksi kalau kolomnya ada, tapi tidak
    # menaikkan syarat minimal.
    alias_columns = [
        alias
        for aliases in COLUMN_ALIASES.values()
        for alias in aliases
    ]

    expected = {
        canonical(col)
        for col in REQUIRED_COLUMNS
        + OPTIONAL_COLUMNS
        + alias_columns
    }

    max_rows = min(
        len(raw_df),
        50,
    )

    best_row = None
    best_score = 0

    for row_idx in range(max_rows):
        values = [
            canonical(value)
            for value in raw_df.iloc[
                row_idx
            ].tolist()
        ]

        score = len(
            set(values) & expected
        )

        if score > best_score:
            best_score = score
            best_row = row_idx

    # Minimal kolom wajib yang harus ditemukan
    # (dikurangi toleransi 2 kolom hilang)
    minimum_required = max(
        1,
        len(REQUIRED_COLUMNS) - 2,
    )

    if best_score < minimum_required:
        return None

    return best_row


def map_standard_columns(df):
    """
    Mapping nama kolom mentah ke nama standar.
    Kalau nama kolom standar tidak ditemukan persis,
    coba cocokkan dengan alias-nya (lihat COLUMN_ALIASES).
    """
    df = df.copy()

    rename_map = {}

    for standard in REQUIRED_COLUMNS + OPTIONAL_COLUMNS:
        target = canonical(standard)

        match_col = None

        # 1. Cocokkan nama kolom standar persis
        for col in df.columns:
            if canonical(col) == target:
                match_col = col
                break

        # 2. Kalau tidak ketemu, coba alias-nya
        if match_col is None:
            for alias in COLUMN_ALIASES.get(
                standard,
                [],
            ):
                alias_target = canonical(alias)

                for col in df.columns:
                    if canonical(col) == alias_target:
                        match_col = col
                        break

                if match_col is not None:
                    break

        if match_col is not None:
            rename_map[match_col] = standard

    return df.rename(
        columns=rename_map
    )


# ============================================================
# FILE READING
# ============================================================

def read_csv_auto(file_bytes):
    """
    Membaca CSV dengan beberapa kemungkinan encoding
    dan mendeteksi header otomatis.
    """
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            raw = pd.read_csv(
                io.BytesIO(file_bytes),
                header=None,
                encoding=encoding,
            )

            header_row = find_header_row(raw)

            if header_row is None:
                continue

            headers = [
                normalize_header(value)
                for value in raw.iloc[
                    header_row
                ].tolist()
            ]

            df = raw.iloc[
                header_row + 1:
            ].copy()

            df.columns = headers

            return df

        except Exception as exc:
            last_error = exc

    raise ValueError(
        f"CSV tidak dapat dibaca. Error: {last_error}"
    )


def load_excel_file(
    file_bytes,
    extension,
):
    """
    Membuka Excel berdasarkan extension.
    """
    if extension == ".xls":
        engine = "xlrd"
    else:
        engine = "openpyxl"

    return pd.ExcelFile(
        io.BytesIO(file_bytes),
        engine=engine,
    )


def load_excel_sheet(
    excel,
    sheet_name,
):
    """
    Membaca sheet Excel tanpa mengasumsikan
    header berada pada baris pertama.
    """
    raw = pd.read_excel(
        excel,
        sheet_name=sheet_name,
        header=None,
    )

    header_row = find_header_row(raw)

    if header_row is None:
        raise ValueError(
            f"Header tidak ditemukan pada sheet "
            f"'{sheet_name}'."
        )

    headers = [
        normalize_header(value)
        for value in raw.iloc[
            header_row
        ].tolist()
    ]

    df = raw.iloc[
        header_row + 1:
    ].copy()

    df.columns = headers

    # Buang kolom kosong / unnamed
    valid_columns = []

    for col in df.columns:
        text = str(col).strip()

        if (
            text != ""
            and not text.lower().startswith(
                "unnamed"
            )
        ):
            valid_columns.append(col)

    df = df[
        valid_columns
    ]

    return df


# ============================================================
# NUMBER PARSING
# ============================================================

def parse_number(value):
    """
    Konversi angka secara fleksibel.

    Contoh:
        1.234
        1,234
        1.234,56
        1,234.56
        Rp 10.000
        (10.000)
    """
    if pd.isna(value):
        return 0.0

    if isinstance(
        value,
        (int, float),
    ):
        return float(value)

    text = str(value).strip()

    if text == "":
        return 0.0

    text = (
        text
        .replace("Rp", "")
        .replace("rp", "")
        .replace(" ", "")
    )

    negative = (
        text.startswith("(")
        and text.endswith(")")
    )

    if negative:
        text = text[1:-1]

    # 1.234,56 atau 1,234.56
    if (
        "." in text
        and "," in text
    ):
        if text.rfind(",") > text.rfind("."):
            # Indonesia:
            # 1.234,56
            text = (
                text
                .replace(".", "")
                .replace(",", ".")
            )
        else:
            # International:
            # 1,234.56
            text = text.replace(
                ",",
                "",
            )

    # Hanya koma
    elif "," in text:
        parts = text.split(",")

        if (
            len(parts) > 1
            and len(parts[-1]) == 3
        ):
            # 1,234
            text = text.replace(
                ",",
                "",
            )
        else:
            # 1,25
            text = text.replace(
                ",",
                ".",
            )

    # Hanya titik
    elif "." in text:
        parts = text.split(".")

        if (
            len(parts) > 1
            and len(parts[-1]) == 3
        ):
            # 1.234
            text = text.replace(
                ".",
                "",
            )

    try:
        number = float(text)
    except Exception:
        number = 0.0

    if negative:
        number = -number

    return number


def clean_numeric(series):
    return series.map(
        parse_number
    ).fillna(0)


# ============================================================
# DATA CLEANING
# ============================================================

def prepare_data(df):
    """
    Cleaning data dan kalkulasi kolom turunan.
    """
    df = df.copy()

    # Hapus baris yang seluruhnya kosong
    df = df.dropna(
        how="all"
    ).reset_index(drop=True)

    # Standardisasi kolom
    df = map_standard_columns(
        df
    )

    missing = [
        col
        for col in REQUIRED_COLUMNS
        if col not in df.columns
    ]

    if missing:
        raise ValueError(
            "Kolom wajib tidak ditemukan: "
            + ", ".join(missing)
        )

    # Text columns
    text_columns = [
        "OpnameNo",
        "Date",
        "Loc",
        "Golongan",
        "Artikel",
        "Barcode",
    ]

    for col in text_columns:
        df[col] = (
            df[col]
            .fillna("")
            .astype(str)
            .str.strip()
        )

    # HM opsional: kalau tidak ada di file, isi 0
    if "HM" not in df.columns:
        df["HM"] = 0.0

    # Numeric columns
    for col in [
        "Qty POS",
        "Qty Fisik",
        "Price",
        "HM",
    ]:
        df[col] = clean_numeric(
            df[col]
        )

    # Remark optional
    if "Remark" not in df.columns:
        df["Remark"] = ""

    df["Remark"] = (
        df["Remark"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    # ========================================================
    # CALCULATIONS
    # ========================================================

    df["Selisih"] = (
        df["Qty Fisik"]
        - df["Qty POS"]
    )

    df["Total"] = (
        df["Selisih"]
        * df["Price"]
    )

    df["Total HM"] = (
        df["Selisih"]
        * df["HM"]
    )

    # Rounding
    df["Qty POS"] = df[
        "Qty POS"
    ].round(6)

    df["Qty Fisik"] = df[
        "Qty Fisik"
    ].round(6)

    df["Selisih"] = df[
        "Selisih"
    ].round(6)

    df["Price"] = df[
        "Price"
    ].round(2)

    df["HM"] = df[
        "HM"
    ].round(2)

    df["Total"] = df[
        "Total"
    ].round(2)

    df["Total HM"] = df[
        "Total HM"
    ].round(2)

    return df


# ============================================================
# BUSINESS LOGIC
# ============================================================

def is_shopbag(row):
    """
    SHOPBAG:
    Golongan atau Artikel mengandung salah satu keyword.
    """
    golongan = str(
        row["Golongan"]
    ).lower()

    artikel = str(
        row["Artikel"]
    ).lower()

    for keyword in SHOPBAG_KEYWORDS:
        keyword = keyword.lower()

        if (
            keyword in golongan
            or keyword in artikel
        ):
            return True

    return False


def categorize(df):
    """
    Prioritas:

    1. SHOPBAG
    2. TERTUKAR
    3. MINUS
    4. PLUS
    5. KLOP
    """
    df = df.copy()

    df["Kategori"] = ""

    # ========================================================
    # 1. SHOPBAG
    # ========================================================

    shopbag_mask = df.apply(
        is_shopbag,
        axis=1,
    )

    df.loc[
        shopbag_mask,
        "Kategori",
    ] = "SHOPBAG"

    # ========================================================
    # 2. TERTUKAR
    # ========================================================

    remaining = (
        df["Kategori"] == ""
    )

    # TERTUKAR: total Selisih per Artikel harus 0
    # (saling menutupi), bukan sekadar ada baris minus
    # dan plus dalam satu Artikel.
    article_totals = (
        df.loc[remaining]
        .groupby("Artikel")[
            "Selisih"
        ]
        .sum()
    )

    exchanged_articles = (
        article_totals[
            article_totals.abs()
            < 1e-6
        ].index
    )

    tertukar_mask = (
        (df["Kategori"] == "")
        & df["Artikel"].isin(
            exchanged_articles
        )
        & (df["Selisih"] != 0)
    )

    df.loc[
        tertukar_mask,
        "Kategori",
    ] = "TERTUKAR"

    # ========================================================
    # 3. MINUS
    # ========================================================

    minus_mask = (
        (df["Kategori"] == "")
        & (df["Selisih"] < 0)
    )

    df.loc[
        minus_mask,
        "Kategori",
    ] = "MINUS"

    # ========================================================
    # 4. PLUS
    # ========================================================

    plus_mask = (
        (df["Kategori"] == "")
        & (df["Selisih"] > 0)
    )

    df.loc[
        plus_mask,
        "Kategori",
    ] = "PLUS"

    # ========================================================
    # 5. KLOP
    # ========================================================

    klop_mask = (
        (df["Kategori"] == "")
        & (df["Selisih"] == 0)
    )

    df.loc[
        klop_mask,
        "Kategori",
    ] = "KLOP"

    # ========================================================
    # REMARK
    # ========================================================
    # Default: nama kategori (Klop/Tertukar/Minus/Plus).
    # Khusus SHOPBAG, ikuti tanda Selisih (Minus/Plus/Klop)
    # karena SHOPBAG sendiri bisa berisi selisih apa saja.

    def get_remark(row):
        category = row["Kategori"]
        selisih = row["Selisih"]

        if category == "SHOPBAG":
            if selisih < 0:
                return "Minus"

            if selisih > 0:
                return "Plus"

            return "Klop"

        return category.capitalize()

    df["Remark"] = df.apply(
        get_remark,
        axis=1,
    )

    return df


# ============================================================
# SUMMARY
# ============================================================

def create_summary(df):
    rows = []

    for category in CATEGORIES:
        part = df[
            df["Kategori"]
            == category
        ]

        rows.append(
            {
                "Kategori": category,
                "Qty Fisik": part[
                    "Qty Fisik"
                ].sum(),
                "Qty POS": part[
                    "Qty POS"
                ].sum(),
                "Selisih": part[
                    "Selisih"
                ].sum(),
                "Nilai Selisih (Rp)": part[
                    "Total"
                ].sum(),
                "HM": part[
                    "HM"
                ].sum(),
                "Total HM": part[
                    "Total HM"
                ].sum(),
            }
        )

    # Grand Total
    rows.append(
        {
            "Kategori": "GRAND TOTAL",
            "Qty Fisik": df[
                "Qty Fisik"
            ].sum(),
            "Qty POS": df[
                "Qty POS"
            ].sum(),
            "Selisih": df[
                "Selisih"
            ].sum(),
            "Nilai Selisih (Rp)": df[
                "Total"
            ].sum(),
            "HM": df[
                "HM"
            ].sum(),
            "Total HM": df[
                "Total HM"
            ].sum(),
        }
    )

    return pd.DataFrame(
        rows
    )


def get_metadata(df):
    def first_value(column):
        values = (
            df[column]
            .dropna()
            .astype(str)
            .str.strip()
        )

        values = values[
            values != ""
        ]

        if len(values):
            return values.iloc[0]

        return ""

    return {
        "Loc": first_value("Loc"),
        "Date": first_value("Date"),
        "OpnameNo": first_value(
            "OpnameNo"
        ),
    }


# ============================================================
# EXCEL STYLE FUNCTIONS
# ============================================================

def apply_number_format(
    cell,
    money=False,
    value=None,
):
    """
    Format angka konsisten dengan section
    NOTE / PERHITUNGAN DEDUCT:

    - money=True  -> pakai prefix "Rp"
                      (mis. -Rp26.404 / Rp0 / Rp500.000)
    - money=False -> angka polos
                      (mis. -12 / 0 / 17)

    Nilai negatif otomatis ditampilkan warna merah.

    Parameter `value` dipakai untuk menentukan warna merah
    kalau cell.value berupa rumus (string, mis. "=E2-F2")
    sehingga tandanya tidak bisa dibaca langsung dari
    cell.value. Kalau tidak diisi, dipakai cell.value.
    """

    if money:
        cell.number_format = (
            '"Rp"#,##0;-"Rp"#,##0;"Rp"0'
        )
    else:
        cell.number_format = "#,##0"

    check_value = (
        cell.value
        if value is None
        else value
    )

    if (
        isinstance(
            check_value,
            (int, float),
        )
        and check_value < 0
    ):
        old_font = cell.font

        cell.font = Font(
            name=old_font.name,
            size=old_font.size,
            bold=old_font.bold,
            italic=old_font.italic,
            color=RED,
        )


def apply_border(
    cell,
    left=None,
    right=None,
    top=None,
    bottom=None,
):
    """
    Mempertahankan border lain ketika
    mengubah satu sisi border.
    """
    old = cell.border

    cell.border = Border(
        left=left or old.left,
        right=right or old.right,
        top=top or old.top,
        bottom=bottom or old.bottom,
    )


def style_summary_header(
    ws,
    row,
    start_col,
    include_hm=True,
):
    headers = [
        "Kategori",
        "Qty Fisik",
        "Qty POS",
        "Selisih",
        "Nilai Selisih (Rp)",
    ]

    if include_hm:
        headers += [
            "HM",
            "Total HM",
        ]

    for offset, header in enumerate(
        headers
    ):
        cell = ws.cell(
            row=row,
            column=start_col + offset,
            value=header,
        )

        is_hm = header in [
            "HM",
            "Total HM",
        ]

        cell.fill = PatternFill(
            "solid",
            fgColor=(
                YELLOW
                if is_hm
                else NAVY
            ),
        )

        cell.font = Font(
            bold=True,
            color=(
                BLACK
                if is_hm
                else WHITE
            ),
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )


def style_detail_header(
    ws,
    row,
    columns=None,
):
    if columns is None:
        columns = DETAIL_COLUMNS

    for col_idx, header in enumerate(
        columns,
        start=1,
    ):
        cell = ws.cell(
            row=row,
            column=col_idx,
            value=header,
        )

        is_hm = header in [
            "HM",
            "Total HM",
        ]

        cell.fill = PatternFill(
            "solid",
            fgColor=(
                YELLOW
                if is_hm
                else NAVY
            ),
        )

        cell.font = Font(
            bold=True,
            color=(
                BLACK
                if is_hm
                else WHITE
            ),
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )


def compute_category_layout(
    df,
    start_row=9,
):
    """
    Hitung posisi baris (data_start_row, data_end_row,
    subtotal_row) untuk tiap kategori di tabel detail,
    SEBELUM tabel detailnya ditulis. Dipakai supaya tabel
    summary di atas bisa merujuk (via rumus) ke baris
    SUBTOTAL / data di tabel detail, sesuai urutan
    penulisan yang dipakai di create_final_excel.
    """

    layout = {}

    row_cursor = start_row

    for category in CATEGORIES:

        part_len = len(
            df[
                df["Kategori"]
                == category
            ]
        )

        title_row = row_cursor
        header_row = row_cursor + 1
        data_start_row = (
            header_row + 1
        )
        subtotal_row = (
            data_start_row
            + part_len
        )

        layout[category] = {
            "title_row": title_row,
            "header_row": header_row,
            "data_start_row": data_start_row,
            "data_end_row": (
                subtotal_row - 1
            ),
            "subtotal_row": subtotal_row,
            "part_len": part_len,
        }

        row_cursor = (
            subtotal_row + 2
        )

    return layout


# ============================================================
# CREATE FINAL EXCEL
# ============================================================

def create_final_excel(
    df,
    metadata,
    summary,
    penalty_percent,
    compensation,
    compensation_note="",
    include_hm=True,
    report_stage=None,
):
    """
    Membuat workbook FINAL.

    report_stage menentukan judul laporan & nama sheet:
      "Sementara" -> LAPORAN ANALISA SEMENTARA / ANALISA SEMENTARA
      "Final"     -> LAPORAN ANALISA FINAL / FINAL
      "Mandiri"   -> LAPORAN ANALISA SO MANDIRI / ANALISA SO MANDIRI
    Kalau None: Final bila include_hm, selain itu Mandiri.

    include_hm=True   -> SO Final (Tim Audit): lengkap
                          dengan kolom HM/Total HM dan
                          section NOTE/PERHITUNGAN DEDUCT.
    include_hm=False  -> SO Mandiri (Toko): kolom HM dan
                          Total HM serta section
                          NOTE/PERHITUNGAN DEDUCT dihilangkan
                          sepenuhnya (bukan cuma disembunyikan).
    """

    # Daftar kolom detail yang benar-benar dipakai di
    # workbook ini (dipakai untuk semua indeks kolom di
    # bawah supaya SO Mandiri otomatis "menggeser" kolom
    # setelah HM/Total HM dihilangkan).
    detail_columns = (
        DETAIL_COLUMNS
        if include_hm
        else [
            col
            for col in DETAIL_COLUMNS
            if col
            not in (
                "HM",
                "Total HM",
            )
        ]
    )

    if report_stage is None:
        report_stage = "Final" if include_hm else "Mandiri"

    stage_titles = {
        "Sementara": (
            "LAPORAN ANALISA SEMENTARA",
            "ANALISA SEMENTARA",
        ),
        "Final": (
            "LAPORAN ANALISA FINAL",
            "FINAL",
        ),
        "Mandiri": (
            "LAPORAN ANALISA SO MANDIRI",
            "ANALISA SO MANDIRI",
        ),
    }

    report_title, sheet_title = stage_titles.get(
        report_stage,
        stage_titles["Final"],
    )

    wb = Workbook()

    ws = wb.active
    ws.title = sheet_title

    ws.sheet_view.showGridLines = False

    # Grouping baris (untuk fitur hide/unhide KATEGORI: KLOP)
    # ditampilkan dengan tombol +/- di BAWAH baris yang
    # dikelompokkan (persis di atas baris SUBTOTAL-nya).
    ws.sheet_properties.outlinePr.summaryBelow = (
        True
    )

    # Posisi baris tiap kategori di tabel detail (dihitung
    # di awal, sebelum ditulis), dipakai supaya tabel
    # summary di atas bisa merujuk via rumus.
    category_layout = compute_category_layout(
        df
    )

    # Huruf kolom Excel untuk setiap nama kolom di
    # DETAIL_COLUMNS, dipakai untuk menulis rumus.
    col_letter = {
        name: get_column_letter(idx)
        for idx, name in enumerate(
            detail_columns,
            start=1,
        )
    }

    # ========================================================
    # HEADER METADATA A-C
    # ========================================================

    ws.merge_cells("A1:C1")

    ws["A1"] = report_title

    ws["A1"].font = Font(
        bold=True,
        size=15,
        color=NAVY,
    )

    ws["A1"].alignment = Alignment(
        vertical="center"
    )

    metadata_fields = [
        (
            "Lokasi / Loc",
            format_location_label(
                metadata["Loc"]
            ),
        ),
        ("Tanggal SO", metadata["Date"]),
        ("No. Opname", metadata["OpnameNo"]),
    ]

    for offset, (label, value) in enumerate(metadata_fields):
        row = 2 + offset

        label_cell = ws.cell(
            row=row,
            column=1,
            value=label,
        )

        # Custom format supaya titik dua (:) di setiap
        # baris sejajar, walau panjang teks labelnya beda.
        label_cell.number_format = (
            '@* ":"'
        )

        label_cell.font = Font(
            bold=True,
        )

        label_cell.alignment = Alignment(
            vertical="center"
        )

        label_cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GRAY,
        )

        # Kolom C ikut diberi warna sebelum di-merge
        # dengan kolom B (nilai)
        ws.cell(
            row=row,
            column=3,
        ).fill = PatternFill(
            "solid",
            fgColor=LIGHT_GRAY,
        )

        # Value dipisah ke kolom B:C (merge biar muat)
        ws.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=3,
        )

        value_cell = ws.cell(
            row=row,
            column=2,
            value=value,
        )

        value_cell.font = Font(
            bold=True,
        )

        value_cell.alignment = Alignment(
            vertical="center"
        )

        value_cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GRAY,
        )

    # ========================================================
    # SUMMARY TABLE D-J
    # ========================================================

    summary_start_col = 4
    summary_header_row = 1

    style_summary_header(
        ws,
        summary_header_row,
        summary_start_col,
        include_hm,
    )

    # Field summary (nama kolom di DataFrame `summary`) dan
    # field detail (nama kolom di DETAIL_COLUMNS) untuk
    # offset 1..6, urutannya harus sama persis.
    summary_record_fields = [
        "Qty Fisik",
        "Qty POS",
        "Selisih",
        "Nilai Selisih (Rp)",
    ]

    summary_detail_fields = [
        "Qty Fisik",
        "Qty POS",
        "Selisih",
        "Total",
    ]

    if include_hm:
        summary_record_fields += [
            "HM",
            "Total HM",
        ]

        summary_detail_fields += [
            "HM",
            "Total HM",
        ]

    summary_row_start = 2
    summary_row_end = (
        summary_row_start
        + len(CATEGORIES)
        - 1
    )

    for excel_row, (
        _,
        record,
    ) in enumerate(
        summary.iterrows(),
        start=2,
    ):
        category_name = record[
            "Kategori"
        ]

        is_grand_total = (
            category_name
            == "GRAND TOTAL"
        )

        # Kolom Kategori (label)
        label_cell = ws.cell(
            row=excel_row,
            column=summary_start_col,
            value=category_name,
        )

        label_cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )

        label_cell.font = Font(
            bold=True
        )

        if is_grand_total:
            label_cell.fill = PatternFill(
                "solid",
                fgColor=LIGHT_GRAY,
            )

        for idx in range(
            len(
                summary_record_fields
            )
        ):
            offset = idx + 1

            record_field = (
                summary_record_fields[
                    idx
                ]
            )

            detail_field = (
                summary_detail_fields[
                    idx
                ]
            )

            raw_value = record[
                record_field
            ]

            if is_grand_total:
                # GRAND TOTAL = jumlah baris kategori
                # di atasnya (dalam tabel summary
                # sendiri).
                summary_letter = (
                    get_column_letter(
                        summary_start_col
                        + offset
                    )
                )

                formula = (
                    f"=SUM({summary_letter}"
                    f"{summary_row_start}:"
                    f"{summary_letter}"
                    f"{summary_row_end})"
                )

            else:
                layout = category_layout[
                    category_name
                ]

                letter = col_letter[
                    detail_field
                ]

                if detail_field == "HM":
                    # Subtotal HM sengaja tidak
                    # ditulis di tabel detail
                    # (rate per unit, bukan nilai
                    # yang dijumlahkan), jadi
                    # dihitung langsung dari
                    # rentang baris datanya.
                    if (
                        layout[
                            "part_len"
                        ]
                        > 0
                    ):
                        formula = (
                            f"=SUM({letter}"
                            f"{layout['data_start_row']}:"
                            f"{letter}"
                            f"{layout['data_end_row']})"
                        )
                    else:
                        formula = (
                            raw_value
                        )
                else:
                    # Field lain sudah punya sel
                    # SUBTOTAL di tabel detail,
                    # tinggal dirujuk.
                    formula = (
                        f"={letter}"
                        f"{layout['subtotal_row']}"
                    )

            cell = ws.cell(
                row=excel_row,
                column=summary_start_col
                + offset,
                value=formula,
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

            apply_number_format(
                cell,
                money=offset
                in [4, 5, 6],
                # Rumus tidak bisa dibaca tandanya
                # langsung, jadi pakai nilai asli
                # hasil hitung Python.
                value=raw_value,
            )

            if is_grand_total:
                is_negative = (
                    isinstance(
                        raw_value,
                        (int, float),
                    )
                    and raw_value < 0
                )

                cell.font = Font(
                    bold=True,
                    color=(
                        RED
                        if is_negative
                        else None
                    ),
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor=LIGHT_GRAY,
                )

                apply_border(
                    cell,
                    bottom=DOUBLE_BLACK,
                )

        if is_grand_total:
            apply_border(
                label_cell,
                bottom=DOUBLE_BLACK,
            )

    # ========================================================
    # DETAIL TABLE
    # ========================================================

    current_row = 9

    for category in CATEGORIES:

        part = df[
            df["Kategori"]
            == category
        ].copy()

        # ----------------------------------------------------
        # CATEGORY TITLE
        # ----------------------------------------------------

        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=len(
                detail_columns
            ),
        )

        title_cell = ws.cell(
            row=current_row,
            column=1,
            value=(
                f"KATEGORI: {category}"
            ),
        )

        title_cell.font = Font(
            bold=True,
            size=12,
            color=NAVY,
        )

        title_cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_BLUE,
        )

        title_cell.alignment = Alignment(
            vertical="center"
        )

        current_row += 1

        # ----------------------------------------------------
        # DETAIL HEADER
        # ----------------------------------------------------

        style_detail_header(
            ws,
            current_row,
            detail_columns,
        )

        current_row += 1

        # ----------------------------------------------------
        # DETAIL DATA
        # ----------------------------------------------------

        data_start_row = current_row

        for _, source in part.iterrows():

            for col_idx, column in enumerate(
                detail_columns,
                start=1,
            ):
                value = source.get(
                    column,
                    "",
                )

                # Jangan tulis NaN
                if pd.isna(value):
                    value = ""

                # Kolom hasil kalkulasi ditulis sebagai
                # rumus Excel (bukan angka statis) supaya
                # bisa diaudit langsung dari file.
                formula = None

                if column == "Selisih":
                    formula = (
                        f"={col_letter['Qty Fisik']}"
                        f"{current_row}-"
                        f"{col_letter['Qty POS']}"
                        f"{current_row}"
                    )

                elif column == "Total":
                    formula = (
                        f"={col_letter['Selisih']}"
                        f"{current_row}*"
                        f"{col_letter['Price']}"
                        f"{current_row}"
                    )

                elif column == "Total HM":
                    formula = (
                        f"={col_letter['Selisih']}"
                        f"{current_row}*"
                        f"{col_letter['HM']}"
                        f"{current_row}"
                    )

                cell = ws.cell(
                    row=current_row,
                    column=col_idx,
                    value=(
                        formula
                        if formula
                        is not None
                        else value
                    ),
                )

                cell.border = Border(
                    left=THIN_GRAY,
                    right=THIN_GRAY,
                    top=THIN_GRAY,
                    bottom=THIN_GRAY,
                )

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=False,
                )

                if column in [
                    "Price",
                    "Qty Fisik",
                    "Qty POS",
                    "Selisih",
                    "Total",
                    "HM",
                    "Total HM",
                ]:
                    apply_number_format(
                        cell,
                        money=column in [
                            "Price",
                            "Total",
                            "HM",
                            "Total HM",
                        ],
                        # Sel Selisih/Total/Total HM
                        # berisi rumus (string), jadi warna
                        # merah ditentukan dari nilai asli
                        # yang sudah dihitung Python.
                        value=value,
                    )

            current_row += 1

        # ----------------------------------------------------
        # HIDE BARIS DATA KHUSUS KATEGORI KLOP
        # (bisa ditampilkan lagi dengan klik tombol +
        # atau unhide manual di Excel)
        # ----------------------------------------------------

        if (
            category == "KLOP"
            and len(part) > 0
        ):
            for hidden_row in range(
                data_start_row,
                current_row,
            ):
                ws.row_dimensions[
                    hidden_row
                ].outlineLevel = 1

                ws.row_dimensions[
                    hidden_row
                ].hidden = True

        # ----------------------------------------------------
        # SUBTOTAL
        # ----------------------------------------------------

        subtotal_row = current_row

        ws.cell(
            row=subtotal_row,
            column=1,
            value="SUBTOTAL",
        )

        subtotal_fields = [
            "Qty Fisik",
            "Qty POS",
            "Selisih",
            "Total",
        ]

        if include_hm:
            subtotal_fields.append(
                "Total HM"
            )

        subtotal_map = {
            detail_columns.index(
                field
            )
            + 1: field
            for field in subtotal_fields
        }

        # Style semua cell subtotal
        for col_idx in range(
            1,
            len(detail_columns) + 1,
        ):
            cell = ws.cell(
                row=subtotal_row,
                column=col_idx,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=LIGHT_GRAY,
            )

            cell.font = Font(
                bold=True
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=DOUBLE_BLACK,
            )

        # Isi subtotal
        has_data_rows = (
            len(part) > 0
        )

        sum_start = data_start_row
        sum_end = current_row - 1

        for col_idx, source_col in subtotal_map.items():

            numeric_value = part[
                source_col
            ].sum()

            letter = col_letter[
                source_col
            ]

            if has_data_rows:
                formula = (
                    f"=SUM({letter}"
                    f"{sum_start}:"
                    f"{letter}"
                    f"{sum_end})"
                )
            else:
                # Tidak ada baris data,
                # tulis 0 langsung
                formula = 0

            cell = ws.cell(
                row=subtotal_row,
                column=col_idx,
                value=formula,
            )

            apply_number_format(
                cell,
                money=source_col in [
                    "Total",
                    "Total HM",
                ],
                # Rumus SUM tidak bisa dibaca
                # tandanya langsung, jadi pakai
                # nilai asli hasil hitung Python.
                value=numeric_value,
            )

        current_row += 2

    # ========================================================
    # NOTE / DEDUCT
    # ========================================================

    # SO Mandiri (include_hm=False): section ini dihilangkan
    # sepenuhnya, bukan cuma disembunyikan.
    if include_hm:
        def apply_rupiah_format(cell):
            """
            Format khusus untuk section NOTE/PERHITUNGAN DEDUCT:
            -Rp26.404 / Rp0 / Rp500.000
            """
            cell.number_format = (
                '"Rp"#,##0;-"Rp"#,##0;"Rp"0'
            )

        note_start = current_row + 1

        ws.merge_cells(
            start_row=note_start,
            start_column=1,
            end_row=note_start,
            end_column=3,
        )

        ws.cell(
            row=note_start,
            column=1,
            value=(
                "NOTE / PERHITUNGAN DEDUCT"
            ),
        )

        ws.cell(
            row=note_start,
            column=1,
        ).font = Font(
            bold=True,
            size=12,
            color=NAVY,
        )

        note_row = note_start + 1

        # --------------------------------------------------------
        # NOTE HEADER
        # --------------------------------------------------------

        note_headers = [
            "Keterangan",
            "Qty Selisih",
            "Total HM",
        ]

        for col_idx, header in enumerate(
            note_headers,
            start=1,
        ):
            is_hm_header = header == "Total HM"

            cell = ws.cell(
                row=note_row,
                column=col_idx,
                value=header,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=(
                    YELLOW
                    if is_hm_header
                    else NAVY
                ),
            )

            cell.font = Font(
                bold=True,
                color=(
                    BLACK
                    if is_hm_header
                    else WHITE
                ),
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

        note_row += 1

        # --------------------------------------------------------
        # CATEGORY VALUES
        # --------------------------------------------------------

        def category_values(category):
            part = df[
                df["Kategori"]
                == category
            ]

            return (
                part["Selisih"].sum(),
                part["Total HM"].sum(),
            )

        shopbag_qty, shopbag_hm = (
            category_values(
                "SHOPBAG"
            )
        )

        minus_qty, minus_hm = (
            category_values(
                "MINUS"
            )
        )

        plus_qty, plus_hm = (
            category_values(
                "PLUS"
            )
        )

        # Label ": nama kategori" -> baris SUBTOTAL kategori
        # terkait di tabel detail (dipakai untuk rumus).
        note_items = [
            (
                "Shopbag",
                "SHOPBAG",
                shopbag_qty,
                shopbag_hm,
            ),
            (
                "Minus",
                "MINUS",
                minus_qty,
                minus_hm,
            ),
            (
                "Plus",
                "PLUS",
                plus_qty,
                plus_hm,
            ),
        ]

        # Simpan nomor baris tiap item note (dipakai untuk
        # rumus "Total Selisih" di bawah).
        note_item_rows = {}

        for (
            label,
            category,
            qty,
            hm,
        ) in note_items:

            note_item_rows[
                category
            ] = note_row

            label_cell = ws.cell(
                row=note_row,
                column=1,
                value=label,
            )

            label_cell.number_format = (
                COLON_LABEL_FORMAT
            )

            label_cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

            layout = category_layout[
                category
            ]

            qty_formula = (
                f"={col_letter['Selisih']}"
                f"{layout['subtotal_row']}"
            )

            hm_formula = (
                f"={col_letter['Total HM']}"
                f"{layout['subtotal_row']}"
            )

            qty_cell = ws.cell(
                row=note_row,
                column=2,
                value=qty_formula,
            )

            hm_cell = ws.cell(
                row=note_row,
                column=3,
                value=hm_formula,
            )

            for cell in (
                qty_cell,
                hm_cell,
            ):
                cell.border = Border(
                    left=THIN_GRAY,
                    right=THIN_GRAY,
                    top=THIN_GRAY,
                    bottom=THIN_GRAY,
                )

            qty_cell.number_format = (
                "#,##0"
            )

            apply_rupiah_format(
                hm_cell
            )

            # Nilai negatif ditandai merah
            # (rumus tidak bisa dibaca tandanya
            # langsung, jadi pakai nilai asli
            # hasil hitung Python).
            if qty < 0:
                qty_cell.font = Font(
                    color=RED
                )

            if hm < 0:
                hm_cell.font = Font(
                    color=RED
                )

            note_row += 1

        # --------------------------------------------------------
        # TOTAL SELISIH
        # --------------------------------------------------------

        total_qty = (
            shopbag_qty
            + minus_qty
            + plus_qty
        )

        total_hm = (
            shopbag_hm
            + minus_hm
            + plus_hm
        )

        total_row = note_row

        shopbag_row = note_item_rows[
            "SHOPBAG"
        ]

        plus_row = note_item_rows[
            "PLUS"
        ]

        ws.cell(
            row=total_row,
            column=1,
            value="Total Selisih",
        )

        total_qty_cell = ws.cell(
            row=total_row,
            column=2,
            value=(
                f"=SUM(B{shopbag_row}:"
                f"B{plus_row})"
            ),
        )

        total_hm_cell = ws.cell(
            row=total_row,
            column=3,
            value=(
                f"=SUM(C{shopbag_row}:"
                f"C{plus_row})"
            ),
        )

        for col_idx in range(1, 4):
            cell = ws.cell(
                row=total_row,
                column=col_idx,
            )

            cell.font = Font(
                bold=True
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=DOUBLE_BLACK,
                bottom=DOUBLE_BLACK,
            )

        total_qty_cell.number_format = (
            "#,##0"
        )

        apply_rupiah_format(
            total_hm_cell
        )

        if total_qty < 0:
            total_qty_cell.font = Font(
                bold=True,
                color=RED,
            )

        if total_hm < 0:
            total_hm_cell.font = Font(
                bold=True,
                color=RED,
            )

        note_row += 1

        # Baris kosong (spasi)
        note_row += 1

        # --------------------------------------------------------
        # TOTAL HM + DENDA
        # --------------------------------------------------------

        denda_value = (
            total_hm
            * penalty_percent
        )

        total_hm_denda = (
            total_hm
            + denda_value
        )

        denda_row = note_row

        ws.cell(
            row=denda_row,
            column=1,
            value="Total HM + Denda",
        ).number_format = COLON_LABEL_FORMAT

        # Persentase denda: input langsung dari user,
        # bukan hasil kalkulasi -> tetap angka biasa.
        denda_percent_cell = ws.cell(
            row=denda_row,
            column=2,
            value=penalty_percent,
        )

        denda_percent_cell.number_format = (
            "0%"
        )

        denda_value_cell = ws.cell(
            row=denda_row,
            column=3,
            value=(
                f"=C{total_row}*"
                f"(1+B{denda_row})"
            ),
        )

        apply_rupiah_format(
            denda_value_cell
        )

        for col_idx in range(1, 4):
            ws.cell(
                row=denda_row,
                column=col_idx,
            ).border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

        if total_hm_denda < 0:
            denda_value_cell.font = Font(
                color=RED,
            )

        note_row += 1

        # --------------------------------------------------------
        # KOMPENSASI
        # --------------------------------------------------------

        kompensasi_row = note_row

        ws.cell(
            row=kompensasi_row,
            column=1,
            value="Kompensasi",
        ).number_format = COLON_LABEL_FORMAT

        # Nilai kompensasi: input langsung dari user,
        # bukan hasil kalkulasi -> tetap angka biasa.
        kompensasi_cell = ws.cell(
            row=kompensasi_row,
            column=3,
            value=compensation,
        )

        apply_rupiah_format(
            kompensasi_cell
        )

        for col_idx in range(1, 4):
            ws.cell(
                row=kompensasi_row,
                column=col_idx,
            ).border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

        note_row += 1

        # Baris kosong (spasi)
        note_row += 1

        # --------------------------------------------------------
        # TOTAL DEDUCT
        # --------------------------------------------------------

        # Kompensasi mengurangi besarnya potongan (Total HM +
        # Denda, yang biasanya minus). Kalau hasilnya masih
        # minus, tampilkan apa adanya; kalau sudah tertutup
        # (>= 0), tampilkan 0 (tidak pernah jadi "untung").
        total_deduct = min(
            0,
            total_hm_denda
            + compensation,
        )

        deduct_row = note_row

        deduct_label_cell = ws.cell(
            row=deduct_row,
            column=1,
            value="Total Deduct",
        )
        deduct_label_cell.font = Font(
            bold=True
        )
        deduct_label_cell.number_format = (
            COLON_LABEL_FORMAT
        )

        for col_idx in range(1, 4):
            ws.cell(
                row=deduct_row,
                column=col_idx,
            ).border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=DOUBLE_BLACK,
                bottom=DOUBLE_BLACK,
            )

        # Deduct wajib kuning
        deduct_cell = ws.cell(
            row=deduct_row,
            column=3,
            value=(
                f"=MIN(0,C{denda_row}"
                f"+C{kompensasi_row})"
            ),
        )

        deduct_cell.fill = PatternFill(
            "solid",
            fgColor=YELLOW,
        )

        deduct_cell.font = Font(
            bold=True,
            color=(
                RED
                if total_deduct < 0
                else BLACK
            ),
        )

        apply_rupiah_format(
            deduct_cell
        )

        note_row += 1

        # --------------------------------------------------------
        # CATATAN KOMPENSASI (teks miring di bawah tabel)
        # --------------------------------------------------------

        if compensation_note:

            catatan_row = note_row + 1

            ws.merge_cells(
                start_row=catatan_row,
                start_column=1,
                end_row=catatan_row,
                end_column=3,
            )

            catatan_cell = ws.cell(
                row=catatan_row,
                column=1,
                value=(
                    f"Catatan Kompensasi: "
                    f"{compensation_note}"
                ),
            )

            catatan_cell.font = Font(
                italic=True,
            )

            catatan_cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            note_row = catatan_row

    # ========================================================
    # GLOBAL EXCEL FORMATTING
    # ========================================================

    ws.freeze_panes = "A11"

    # --------------------------------------------------------
    # Auto-fit lebar kolom (dihitung dari isi tiap kolom)
    # --------------------------------------------------------

    def estimate_formatted_length(
        raw_value,
        money=False,
    ):
        """
        Perkiraan panjang teks angka setelah diformat
        (dipakai untuk kolom yang isinya rumus, karena
        openpyxl tidak bisa membaca hasil rumus).
        """
        try:
            number = float(raw_value)
        except (TypeError, ValueError):
            return len(
                str(raw_value)
            )

        text = f"{abs(number):,.0f}"

        if money:
            text = "Rp" + text

        if number < 0:
            text = "-" + text

        return len(text)

    MIN_COLUMN_WIDTH = 10
    MAX_COLUMN_WIDTH = 40
    COLUMN_PADDING = 2

    # Sel yang jadi "anchor" dari merge lebih dari 1 kolom
    # dilewati dari perhitungan (teksnya tidak perlu muat
    # dalam satu kolom saja, mis. judul & catatan panjang).
    wide_merge_anchors = set()

    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.max_col
            > merged_range.min_col
        ):
            wide_merge_anchors.add(
                (
                    merged_range.min_row,
                    merged_range.min_col,
                )
            )

    max_lengths = {}

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue

            if (
                cell.row,
                cell.column,
            ) in wide_merge_anchors:
                continue

            text = str(cell.value)

            # Sel berisi rumus: panjang teks rumusnya
            # tidak merepresentasikan tampilan akhir,
            # jadi dilewati (diestimasi terpisah di bawah).
            if text.startswith("="):
                continue

            col = cell.column_letter

            max_lengths[col] = max(
                max_lengths.get(
                    col, 0
                ),
                len(text),
            )

    # Kolom hasil rumus di tabel detail (Selisih, Total,
    # Total HM) diestimasi dari nilai asli di dataframe.
    formula_columns = {
        col_letter["Selisih"]: (
            df["Selisih"],
            False,
        ),
        col_letter["Total"]: (
            df["Total"],
            True,
        ),
    }

    if include_hm:
        formula_columns[
            col_letter["Total HM"]
        ] = (
            df["Total HM"],
            True,
        )

    for col, (
        series,
        money,
    ) in formula_columns.items():

        if len(series) == 0:
            continue

        estimated = max(
            estimate_formatted_length(
                v,
                money=money,
            )
            for v in series
        )

        max_lengths[col] = max(
            max_lengths.get(
                col, 0
            ),
            estimated,
        )

    for col, length in max_lengths.items():
        width = min(
            max(
                length
                + COLUMN_PADDING,
                MIN_COLUMN_WIDTH,
            ),
            MAX_COLUMN_WIDTH,
        )

        ws.column_dimensions[
            col
        ].width = width

    # --------------------------------------------------------
    # Row height
    # --------------------------------------------------------

    ws.row_dimensions[1].height = 26

    # --------------------------------------------------------
    # Sembunyikan kolom HM (harga satuan mentah). Kolomnya
    # tetap ada (dipakai rumus Total HM), cuma disembunyikan
    # dari tampilan supaya rapi.
    # --------------------------------------------------------

    if include_hm and "HM" in col_letter:
        ws.column_dimensions[
            col_letter["HM"]
        ].hidden = True

    # --------------------------------------------------------
    # General alignment
    # --------------------------------------------------------

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                # Jangan mengubah alignment note yang sudah wrap
                if cell.alignment.wrap_text is None:
                    cell.alignment = Alignment(
                        vertical="center"
                    )

    # ========================================================
    # OUTPUT BYTES
    # ========================================================

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# CREATE BERITA ACARA (DOCX)
# ============================================================

BA_NAVY = RGBColor(0x1F, 0x49, 0x7D)
BA_YELLOW = "FFFF00"
BA_HEADER_BLUE = "1F497D"
BA_RED = RGBColor(0xFF, 0x00, 0x00)
BA_BLACK = RGBColor(0x00, 0x00, 0x00)


def rupiah_text(value):
    """Format angka jadi teks Rupiah gaya Indonesia,
    mis. -506000 -> '-Rp506.000', 0 -> 'Rp0'."""

    value = float(value or 0)
    sign = "-" if value < 0 else ""
    text = f"{abs(value):,.0f}".replace(",", ".")

    return f"{sign}Rp{text}"


def qty_text(value):
    value = float(value or 0)
    sign = "-" if value < 0 else ""
    text = f"{abs(value):,.0f}".replace(",", ".")

    return f"{sign}{text}"


# ============================================================
# EMAIL ANALISA SO
# ============================================================

INDONESIAN_MONTHS = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]

DEFAULT_EMAIL_CLOSING = (
    "Tolong kroscek dan report kembali segera selisih "
    "plus atau minusnya serta mengirimkan bukti cek "
    "fisik dengan video via WA IC."
)


def format_tanggal_id(value):
    """
    Ubah tanggal apa pun dari file (2026-09-19, 19/09/2026,
    19-Sep-2026, dst) jadi "19 September 2026".
    Kalau sudah berupa nama bulan Indonesia atau tidak bisa
    dibaca, dikembalikan apa adanya.
    """

    text = str(value or "").strip()

    if not text:
        return ""

    lowered = text.lower()

    if any(
        month.lower() in lowered
        for month in INDONESIAN_MONTHS
    ):
        return text

    # Format tahun-dulu (2026-09-02, 2026/09/02) SELALU
    # tahun-bulan-hari. Kalau dipaksa dayfirst, pandas bisa
    # membalik bulan & hari (2026-09-02 -> 9 Februari).
    year_first = bool(
        re.match(r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", text)
    )

    with warnings.catch_warnings():

        warnings.simplefilter("ignore")

        parsed = pd.to_datetime(
            text,
            dayfirst=not year_first,
            errors="coerce",
        )

        if pd.isna(parsed):
            parsed = pd.to_datetime(
                text,
                errors="coerce",
            )

    if pd.isna(parsed):
        return text

    return (
        f"{parsed.day} "
        f"{INDONESIAN_MONTHS[parsed.month - 1]} "
        f"{parsed.year}"
    )


def create_analisa_email(
    summary,
    metadata,
    so_stage,
    tanggal_so_text="",
    closing_text=DEFAULT_EMAIL_CLOSING,
    sender_name="",
    penalty_percent=0.0,
    compensation=0.0,
    compensation_note="",
):
    """
    Bangun subject + body HTML email "Analisa SO".

    so_stage:
      "Mandiri"   -> tabel rekap 5 kolom + kalimat penutup.
      "Sementara" -> tabel rekap 5 kolom (tanpa kalimat penutup).
      "Final"     -> tabel rekap + kolom Total HM, lalu blok
                     NOTE / PERHITUNGAN DEDUCT (sama dengan
                     Berita Acara) dan Catatan Kompensasi
                     (tanpa kalimat penutup).

    Semua style ditulis inline supaya formatnya tetap utuh
    saat di-paste ke Gmail / Outlook.
    Return: (subject, html_body)
    """

    esc = html_escape

    with_deduct = so_stage == "Final"

    loc_code = str(metadata.get("Loc") or "").strip()
    initial = LOCATION_INITIALS.get(loc_code, loc_code) or "-"

    tanggal = format_tanggal_id(
        tanggal_so_text or metadata.get("Date")
    ) or "-"

    if so_stage == "Final":
        subject = f"FINAL ANALISA SO IC {initial} {tanggal}"
        intro = (
            f"Berikut terlampir Final Analisa SO IC "
            f"{initial} {tanggal} :"
        )

    elif so_stage == "Sementara":
        subject = (
            f"ANALISA SEMENTARA SO IC {initial} {tanggal}"
        )
        intro = (
            f"Berikut terlampir Analisa Sementara SO IC "
            f"{initial} {tanggal} :"
        )

    else:
        subject = f"ANALISA SO MANDIRI {initial} {tanggal}"
        intro = (
            f"Berikut terlampir Analisa SO Mandiri "
            f"{initial} {tanggal} :"
        )

    subject = subject.upper()

    # ---------------- style dasar ----------------

    border = "1px solid #808080"
    RED_HEX = "#FF0000"
    YELLOW_HEX = "#FFFF00"

    font = (
        "font-family:Arial,Helvetica,sans-serif;"
        "font-size:13px;color:#000000;"
    )

    def th(text, yellow=False, align="center", width=None):

        bg = YELLOW_HEX if yellow else f"#{NAVY}"
        fg = "#000000" if yellow else "#FFFFFF"

        style = (
            f"background-color:{bg};color:{fg};"
            f"font-weight:bold;text-align:{align};"
            f"border:{border};padding:4px 8px;"
        )

        if width:
            style += f"width:{width}px;"

        return f'<th style="{style}">{esc(text)}</th>'

    def td(
        text,
        align="right",
        bold=False,
        red=False,
        bg=None,
        heavy_top=False,
        width=None,
        raw=False,
    ):

        style = (
            f"border:{border};padding:3px 8px;"
            f"text-align:{align};"
        )

        if bg:
            style += f"background-color:{bg};"

        if heavy_top:
            style += "border-top:2px solid #000000;"

        if bold:
            style += "font-weight:bold;"

        if red:
            style += f"color:{RED_HEX};"

        if width:
            style += f"width:{width}px;"

        content = text if raw else esc(text)

        return f'<td style="{style}">{content}</td>'

    def colon_label(text, colon=True, bold=False):
        """Label dengan titik dua rata kanan di dalam sel."""

        weight = "font-weight:bold;" if bold else ""

        if not colon:
            return (
                f'<span style="{font}{weight}">'
                f"{esc(text)}</span>"
            )

        return (
            f'<table style="width:100%;'
            f'border-collapse:collapse;{font}">'
            "<tr>"
            f'<td style="padding:0;text-align:left;'
            f'white-space:nowrap;{font}{weight}">'
            f"{esc(text)}</td>"
            f'<td style="padding:0;text-align:right;'
            f'width:10px;{font}{weight}">:</td>'
            "</tr></table>"
        )

    # ---------------- tabel rekap ----------------

    headers = [
        th("Kategori"),
        th("Qty Fisik"),
        th("Qty POS"),
        th("Selisih"),
        th("Nilai Selisih (Rp)"),
    ]

    if with_deduct:
        headers.append(th("Total HM", yellow=True))

    rows_html = ""

    for _, row in summary.iterrows():

        is_total = row["Kategori"] == "GRAND TOTAL"

        opts = {
            "bold": is_total,
            "bg": "#F2F2F2" if is_total else None,
            "heavy_top": is_total,
        }

        selisih = float(row["Selisih"] or 0)
        nilai = float(row["Nilai Selisih (Rp)"] or 0)

        # Merah hanya untuk nilai negatif (minus).
        cells = [
            td(
                row["Kategori"],
                align="left",
                **{**opts, "bold": True},
            ),
            td(qty_text(row["Qty Fisik"]), **opts),
            td(qty_text(row["Qty POS"]), **opts),
            td(
                qty_text(selisih),
                red=selisih < 0,
                **opts,
            ),
            td(
                rupiah_text(nilai),
                red=nilai < 0,
                **opts,
            ),
        ]

        if with_deduct:

            total_hm_row = float(row["Total HM"] or 0)

            cells.append(
                td(
                    rupiah_text(total_hm_row),
                    red=total_hm_row < 0,
                    **opts,
                )
            )

        rows_html += "<tr>" + "".join(cells) + "</tr>"

    table_html = (
        f'<table style="border-collapse:collapse;{font}">'
        f'<thead><tr>{"".join(headers)}</tr></thead>'
        f"<tbody>{rows_html}</tbody>"
        "</table>"
    )

    # ---------------- body ----------------

    gap = 'style="margin:0 0 14px 0;"'

    body = (
        f'<div style="{font}">'
        f"<p {gap}>Dear {esc(initial)},</p>"
        f"<p {gap}>{esc(intro)}</p>"
        f"<div {gap}>{table_html}</div>"
    )

    # ---------------- NOTE / PERHITUNGAN DEDUCT (Final) ------

    if with_deduct:

        def pick(category):

            match = summary[summary["Kategori"] == category]

            if len(match):
                item = match.iloc[0]
                return (
                    float(item["Selisih"] or 0),
                    float(item["Total HM"] or 0),
                )

            return 0.0, 0.0

        shop_q, shop_hm = pick("SHOPBAG")
        minus_q, minus_hm = pick("MINUS")
        plus_q, plus_hm = pick("PLUS")

        total_q = shop_q + minus_q + plus_q
        total_hm = shop_hm + minus_hm + plus_hm

        total_hm_denda = total_hm * (
            1 + penalty_percent / 100
        )
        total_deduct = min(
            0, total_hm_denda + compensation
        )

        # Lebar kolom disamakan di semua tabel deduct
        # supaya sejajar rapi.
        W1, W2, W3 = 160, 85, 100

        table_style = (
            f"border-collapse:collapse;{font}"
        )

        note_rows = ""

        for label, qty, hm in [
            ("Shopbag", shop_q, shop_hm),
            ("Minus", minus_q, minus_hm),
            ("Plus", plus_q, plus_hm),
        ]:
            note_rows += (
                "<tr>"
                + td(
                    colon_label(label),
                    align="left",
                    raw=True,
                    width=W1,
                )
                + td(
                    qty_text(qty),
                    red=qty < 0,
                    width=W2,
                )
                + td(
                    rupiah_text(hm),
                    red=hm < 0,
                    width=W3,
                )
                + "</tr>"
            )

        note_rows += (
            "<tr>"
            + td(
                "Total Selisih",
                align="left",
                bold=True,
                heavy_top=True,
                width=W1,
            )
            + td(
                qty_text(total_q),
                bold=True,
                red=total_q < 0,
                heavy_top=True,
                width=W2,
            )
            + td(
                rupiah_text(total_hm),
                bold=True,
                red=total_hm < 0,
                heavy_top=True,
                width=W3,
            )
            + "</tr>"
        )

        body += (
            '<p style="margin:0 0 6px 0;font-weight:bold;'
            f'font-size:14px;color:#{NAVY};">'
            "NOTE / PERHITUNGAN DEDUCT</p>"
            f'<table style="{table_style}">'
            "<thead><tr>"
            + th("Keterangan", align="left", width=W1)
            + th("Qty Selisih", width=W2)
            + th("Total HM", yellow=True, width=W3)
            + "</tr></thead>"
            f"<tbody>{note_rows}</tbody></table>"
        )

        body += (
            f'<table style="{table_style}margin-top:14px;">'
            "<tbody>"
            "<tr>"
            + td(
                colon_label("Total HM + Denda"),
                align="left",
                raw=True,
                width=W1,
            )
            + td(f"{penalty_percent:.0f}%", width=W2)
            + td(
                rupiah_text(total_hm_denda),
                red=total_hm_denda < 0,
                width=W3,
            )
            + "</tr><tr>"
            + td(
                colon_label("Kompensasi", colon=False),
                align="left",
                raw=True,
                width=W1,
            )
            + td("", width=W2)
            + td(rupiah_text(compensation), width=W3)
            + "</tr></tbody></table>"
        )

        body += (
            f'<table style="{table_style}margin-top:14px;">'
            "<tbody><tr>"
            + td(
                colon_label("Total Deduct", bold=True),
                align="left",
                raw=True,
                width=W1,
                heavy_top=True,
            )
            + td("", width=W2, heavy_top=True)
            + td(
                rupiah_text(total_deduct),
                bold=True,
                red=total_deduct < 0,
                bg=YELLOW_HEX,
                width=W3,
                heavy_top=True,
            )
            + "</tr></tbody></table>"
        )

        if compensation_note and compensation_note.strip():
            note_html = esc(
                compensation_note.strip()
            ).replace("\n", "<br>")

            body += (
                '<p style="margin:14px 0 0 0;">'
                f"<i>Catatan Kompensasi: {note_html}</i></p>"
            )

        body += '<div style="height:14px;"></div>'

    # ---------------- kalimat penutup (hanya Mandiri) --------

    closing_html = esc(
        (closing_text or "").strip()
    ).replace("\n", "<br>")

    if closing_html and so_stage == "Mandiri":
        body += (
            f"<p {gap}><i>{closing_html}</i></p>"
        )

    body += (
        f'<p style="margin:0;">Thanks,<br>'
        f"{esc(sender_name)}</p>"
        "</div>"
    )

    return subject, body


def render_email_preview(email_html, height=560):
    """
    Tampilkan preview email + tombol "Copy isi email".
    Hasil copy berupa rich text (tabel lengkap dengan warna),
    tinggal Ctrl+V di kolom compose Gmail / Outlook.
    """

    page = """
    <div style="font-family:Arial,sans-serif;">
      <button id="copyBtn" style="
        background:#1A73E8;color:#fff;border:none;
        padding:8px 16px;border-radius:6px;cursor:pointer;
        font-size:14px;">📋 Copy isi email</button>
      <span id="msg" style="margin-left:10px;font-size:13px;
        color:#188038;"></span>
      <div id="emailBody" style="background:#fff;color:#000;
        border:1px solid #ddd;border-radius:6px;padding:16px;
        margin-top:12px;">__EMAIL_HTML__</div>
    </div>
    <script>
    const msg = document.getElementById('msg');
    const el = document.getElementById('emailBody');

    function done(ok) {
      msg.style.color = ok ? '#188038' : '#D93025';
      msg.textContent = ok
        ? '✔ Tersalin — paste (Ctrl+V) di email'
        : 'Gagal otomatis, blok isi email lalu Ctrl+C';
    }

    function fallbackCopy() {
      const range = document.createRange();
      range.selectNodeContents(el);
      const sel = window.getSelection();
      sel.removeAllRanges();
      sel.addRange(range);
      let ok = false;
      try { ok = document.execCommand('copy'); } catch (e) {}
      sel.removeAllRanges();
      done(ok);
    }

    document.getElementById('copyBtn')
      .addEventListener('click', async () => {
        try {
          if (navigator.clipboard && window.ClipboardItem) {
            await navigator.clipboard.write([
              new ClipboardItem({
                'text/html': new Blob([el.innerHTML],
                  {type: 'text/html'}),
                'text/plain': new Blob([el.innerText],
                  {type: 'text/plain'})
              })
            ]);
            done(true);
            return;
          }
        } catch (e) {}
        fallbackCopy();
      });
    </script>
    """.replace("__EMAIL_HTML__", email_html)

    components.html(
        page,
        height=height,
        scrolling=True,
    )


def _set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")

    for tag, value in (
        ("top", top),
        ("bottom", bottom),
        ("start", left),
        ("end", right),
    ):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)

    tcPr.append(mar)


def _set_cell_border(cell, **kwargs):
    """kwargs: top/bottom/left/right, tiap-tiap dict berisi
    val ('single'/'double'/'none'), sz, color."""

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))

    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)

    for edge in ("top", "left", "bottom", "right"):
        edge_data = kwargs.get(edge)

        if not edge_data:
            continue

        tag = f"w:{edge}"
        element = tcBorders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)

        element.set(qn("w:val"), edge_data.get("val", "single"))
        element.set(qn("w:sz"), str(edge_data.get("sz", 4)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), edge_data.get("color", "000000"))


def _set_row_double_border(table, row_idx, top=True, bottom=True, sz=6):
    for cell in table.rows[row_idx].cells:
        kwargs = {}

        if top:
            kwargs["top"] = {
                "val": "double",
                "sz": sz,
                "color": "000000",
            }

        if bottom:
            kwargs["bottom"] = {
                "val": "double",
                "sz": sz,
                "color": "000000",
            }

        _set_cell_border(cell, **kwargs)


def _apply_thin_gray_borders(table, color="B7B7B7", sz=4):
    """Kasih border tipis abu-abu (mirip THIN_GRAY di Excel)
    ke semua sel tabel, supaya tampilan tabel di Word sama
    persis seperti tabel di Excel. Dipanggil sebelum override
    border ganda (_set_row_double_border) supaya border
    ganda tetap menang di baris/kolom yang relevan."""

    edge = {
        "val": "single",
        "sz": sz,
        "color": color,
    }

    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(
                cell,
                top=edge,
                bottom=edge,
                left=edge,
                right=edge,
            )


def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "none")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)

    tblPr.append(borders)


def _write_cell(
    cell,
    text,
    bold=False,
    italic=False,
    size=10,
    color=None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    valign=WD_ALIGN_VERTICAL.CENTER,
    shading=None,
):
    cell.text = ""
    cell.vertical_alignment = valign

    paragraph = cell.paragraphs[0]
    paragraph.alignment = align

    run = paragraph.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)

    if color is not None:
        run.font.color.rgb = color

    if shading:
        _set_cell_shading(cell, shading)

    _set_cell_margins(cell)

    return run


def _add_paragraph_border(paragraph, sz=18, color="000000"):
    """Kasih border bawah pada paragraph (dipakai untuk garis
    di bawah letterhead)."""

    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)

    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_box_border(paragraph, sz=12, color="000000"):
    """Kasih border sekeliling paragraph (dipakai untuk kotak
    judul BERITA ACARA)."""

    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")

    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(sz))
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), color)
        pBdr.append(node)

    pPr.append(pBdr)


def _set_table_fixed_layout(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def _autofit_table_columns(
    table,
    column_texts,
    min_cm=1.4,
    max_cm=6.5,
    char_width_cm=0.2,
    padding_cm=0.45,
):
    """
    Hitung lebar tiap kolom dari panjang teks terpanjang di
    kolom itu (header + isi), mirip fitur "AutoFit Contents"
    di Word. python-docx tidak bisa mengukur lebar teks
    sungguhan (itu dihitung Word saat dibuka), jadi ini
    estimasi berbasis jumlah karakter, lalu lebarnya DIKUNCI
    (tblLayout fixed) supaya tampil konsisten di Word,
    LibreOffice, maupun Google Docs.

    column_texts: list berisi satu list-of-str per kolom,
    semua teks yang muncul di kolom tsb (termasuk header).
    """

    widths_cm = []

    for texts in column_texts:
        max_len = max(
            (len(str(t)) for t in texts),
            default=0,
        )

        width_cm = min(
            max(
                max_len * char_width_cm
                + padding_cm,
                min_cm,
            ),
            max_cm,
        )

        widths_cm.append(width_cm)

    for row in table.rows:
        for idx, cell in enumerate(
            row.cells
        ):
            cell.width = Cm(
                widths_cm[idx]
            )

    for idx, column in enumerate(
        table.columns
    ):
        column.width = Cm(
            widths_cm[idx]
        )

    _set_table_fixed_layout(table)

    return widths_cm


def _add_info_table(doc, rows, label_width_cm=3.0):
    """Tabel tanpa border 2 kolom (label, nilai) supaya tanda
    titik-dua semuanya sejajar rapi, sama seperti dokumen asli.

    value_runs: list of (text, bold) atau (text, bold, highlight).
    highlight=True menandai teks dengan stabilo kuning (mis.
    nama Loc), sama seperti blok kuning di versi Excel."""

    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False

    for r, (label, value_runs) in enumerate(rows):
        table.cell(r, 0).width = Cm(label_width_cm)
        table.cell(r, 1).width = Cm(13)

        label_cell = table.cell(r, 0)
        _write_cell(
            label_cell,
            label,
            size=10.5,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )
        label_cell.paragraphs[0].paragraph_format.space_after = Pt(0)

        value_cell = table.cell(r, 1)
        value_cell.text = ""
        p = value_cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)

        colon_run = p.add_run(": ")
        colon_run.font.size = Pt(10.5)

        for value_run in value_runs:
            text, bold = value_run[0], value_run[1]
            highlight = (
                value_run[2]
                if len(value_run) > 2
                else False
            )

            run = p.add_run(text)
            run.bold = bold
            run.font.size = Pt(10.5)

            if highlight:
                run.font.highlight_color = (
                    WD_COLOR_INDEX.YELLOW
                )

        _set_cell_margins(value_cell, top=15, bottom=15)
        _set_cell_margins(label_cell, top=15, bottom=15)

    _set_table_fixed_layout(table)
    _remove_table_borders(table)

    return table


def create_berita_acara_docx(
    metadata,
    summary,
    is_final,
    penalty_percent,
    compensation,
    compensation_note,
    company_name,
    kepada,
    perihal,
    no_adjustment,
    tanggal_so_text,
    kota,
    tanggal_ba_text,
    dibuat_oleh,
    mengetahui_list,
    menyetujui_list,
):
    """Membuat dokumen Berita Acara (.docx) sesuai format
    standar (letterhead, kotak judul, tabel departemen,
    info SO, tabel kategori, tabel keterangan/deduct untuk
    SO Final, dan blok tanda tangan)."""

    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_before = Pt(0)
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.0

    loc_label = format_location_full(
        metadata.get("Loc", "")
    )

    grand = summary[
        summary["Kategori"] == "GRAND TOTAL"
    ].iloc[0]

    total_qty = grand["Selisih"]
    total_harga = grand["Nilai Selisih (Rp)"]

    # --------------------------------------------------------
    # LETTERHEAD
    # --------------------------------------------------------

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER

    try:
        logo_stream = io.BytesIO(
            base64.b64decode(HARDWARE_LOGO_B64)
        )
        run = head.add_run()
        run.add_picture(logo_stream, width=Cm(6.5))
    except Exception:
        # Fallback ke teks kalau logo gagal dimuat
        run = head.add_run(company_name.upper())
        run.bold = True
        run.font.size = Pt(30)
        run.font.name = "Arial"

    _add_paragraph_border(head)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --------------------------------------------------------
    # TITLE BOX
    # --------------------------------------------------------

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(6)
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run("BERITA ACARA")
    title_run.bold = True
    title_run.font.size = Pt(16)
    _add_box_border(title_p)

    doc.add_paragraph()

    # --------------------------------------------------------
    # DEPARTMENT TABLE (statis)
    # --------------------------------------------------------

    dept_rows = [
        [
            "Human Resource",
            "Operational",
            "Fin & Accounting",
            "Markom",
            "Merchandiser",
        ],
        [
            "Ware House",
            "IT",
            "Inventory Control",
            "VM",
            "Fashion Design",
        ],
    ]

    dept_table = doc.add_table(rows=2, cols=5)
    dept_table.style = "Table Grid"
    dept_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r, row_values in enumerate(dept_rows):
        for c, value in enumerate(row_values):
            _write_cell(
                dept_table.cell(r, c),
                value,
                bold=True,
                size=9.5,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )

    doc.add_paragraph()

    # --------------------------------------------------------
    # INFO BLOCK: Kepada / Perihal / Loc
    # --------------------------------------------------------

    _add_info_table(
        doc,
        [
            ("Kepada", [(kepada, False)]),
            ("Perihal", [(perihal, False)]),
            ("Loc", [(loc_label, True, True)]),
        ],
    )

    doc.add_paragraph()

    hormat = doc.add_paragraph("Dengan hormat ,")
    hormat.paragraph_format.space_after = Pt(6)

    body = doc.add_paragraph()
    body.paragraph_format.space_after = Pt(6)
    body.add_run(
        "Dengan dibuatnya berita acara ini, kami dari Inv. "
        "Department mengajukan untuk melakukan proses posting "
        "adjustment SO IC "
    ).font.size = Pt(10.5)
    bold_loc = body.add_run(f"{loc_label} ")
    bold_loc.bold = True
    bold_loc.font.size = Pt(10.5)
    body.add_run(
        f"tanggal {tanggal_so_text}."
    ).font.size = Pt(10.5)

    doc.add_paragraph()

    _add_info_table(
        doc,
        [
            ("No ADJUSTMENT", [(no_adjustment, True)]),
            ("Total Qty", [(f"{qty_text(total_qty)} Pcs", False)]),
            ("Total Harga", [(f"Rp. {qty_text(total_harga)}", False)]),
        ],
    )

    doc.add_paragraph()

    # --------------------------------------------------------
    # CATEGORY TABLE
    # --------------------------------------------------------

    cat_headers = [
        "Kategori",
        "Qty Fisik",
        "Qty POS",
        "Selisih",
        "Nilai Selisih (Rp)",
    ]

    if is_final:
        cat_headers.append("Total HM")

    cat_table = doc.add_table(
        rows=1 + len(summary), cols=len(cat_headers)
    )
    cat_table.style = "Table Grid"

    cat_column_texts = [
        [header] for header in cat_headers
    ]

    for c, header in enumerate(cat_headers):
        is_hm_header = header == "Total HM"

        _write_cell(
            cat_table.cell(0, c),
            header,
            bold=True,
            size=9.5,
            color=(
                BA_BLACK
                if is_hm_header
                else RGBColor(0xFF, 0xFF, 0xFF)
            ),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            shading=(
                BA_YELLOW
                if is_hm_header
                else BA_HEADER_BLUE
            ),
        )

    for r, (_, row) in enumerate(summary.iterrows(), start=1):
        is_grand = row["Kategori"] == "GRAND TOTAL"

        values = [
            row["Kategori"],
            qty_text(row["Qty Fisik"]),
            qty_text(row["Qty POS"]),
            qty_text(row["Selisih"]),
            rupiah_text(row["Nilai Selisih (Rp)"]),
        ]

        if is_final:
            values.append(rupiah_text(row["Total HM"]))

        for c, value in enumerate(values):
            is_negative = (
                c >= 3
                and isinstance(value, str)
                and value.startswith("-")
            )

            _write_cell(
                cat_table.cell(r, c),
                value,
                bold=is_grand,
                size=9.5,
                color=BA_RED if is_negative else None,
                align=(
                    WD_ALIGN_PARAGRAPH.LEFT
                    if c == 0
                    else WD_ALIGN_PARAGRAPH.RIGHT
                ),
            )

            cat_column_texts[c].append(str(value))

    _apply_thin_gray_borders(cat_table)
    _set_row_double_border(cat_table, len(summary))

    _autofit_table_columns(
        cat_table, cat_column_texts
    )

    doc.add_paragraph()

    # --------------------------------------------------------
    # KETERANGAN / DEDUCT TABLE (SO Final saja)
    # --------------------------------------------------------

    if is_final:

        def get_row(kategori):
            match = summary[summary["Kategori"] == kategori]
            if len(match):
                return match.iloc[0]
            return {"Selisih": 0, "Total HM": 0}

        shopbag = get_row("SHOPBAG")
        minus = get_row("MINUS")
        plus = get_row("PLUS")

        total_selisih_qty = (
            shopbag["Selisih"] + minus["Selisih"] + plus["Selisih"]
        )
        total_selisih_hm = (
            shopbag["Total HM"] + minus["Total HM"] + plus["Total HM"]
        )

        denda_value = total_selisih_hm * (penalty_percent / 100)
        total_hm_denda = total_selisih_hm + denda_value
        total_deduct = min(0, total_hm_denda + compensation)

        note_rows = [
            ("Shopbag :", shopbag["Selisih"], shopbag["Total HM"], False),
            ("Minus :", minus["Selisih"], minus["Total HM"], False),
            ("Plus :", plus["Selisih"], plus["Total HM"], False),
            (
                "Total Selisih",
                total_selisih_qty,
                total_selisih_hm,
                True,
            ),
        ]

        note_table = doc.add_table(
            rows=1 + len(note_rows), cols=3
        )
        note_table.style = "Table Grid"

        note_headers = ["Keterangan", "Qty Selisih", "Total HM"]
        note_column_texts = [
            [header] for header in note_headers
        ]

        for c, header in enumerate(note_headers):
            is_hm_header = header == "Total HM"

            _write_cell(
                note_table.cell(0, c),
                header,
                bold=True,
                size=9.5,
                color=(
                    BA_BLACK
                    if is_hm_header
                    else RGBColor(0xFF, 0xFF, 0xFF)
                ),
                align=(
                    WD_ALIGN_PARAGRAPH.LEFT
                    if c == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                ),
                shading=(
                    BA_YELLOW
                    if is_hm_header
                    else BA_HEADER_BLUE
                ),
            )

        for r, (label, qty, hm, is_total) in enumerate(
            note_rows, start=1
        ):
            row_values = [label, qty_text(qty), rupiah_text(hm)]

            for c, value in enumerate(row_values):
                is_negative = c >= 1 and str(value).startswith("-")

                _write_cell(
                    note_table.cell(r, c),
                    value,
                    bold=is_total,
                    size=9.5,
                    color=BA_RED if is_negative else None,
                    align=(
                        WD_ALIGN_PARAGRAPH.LEFT
                        if c == 0
                        else WD_ALIGN_PARAGRAPH.RIGHT
                    ),
                )

                note_column_texts[c].append(str(value))

        _apply_thin_gray_borders(note_table)
        _set_row_double_border(note_table, len(note_rows))

        _autofit_table_columns(
            note_table, note_column_texts
        )

        doc.add_paragraph()

        deduct_table = doc.add_table(rows=2, cols=3)
        deduct_table.style = "Table Grid"

        deduct_row0_texts = [
            "Total HM + Denda :",
            f"{penalty_percent:.0f}%",
            rupiah_text(total_hm_denda),
        ]

        deduct_row1_texts = [
            "Kompensasi :",
            "",
            rupiah_text(compensation),
        ]

        _write_cell(
            deduct_table.cell(0, 0),
            "Total HM + Denda :",
            size=9.5,
        )
        _write_cell(
            deduct_table.cell(0, 1),
            f"{penalty_percent:.0f}%",
            size=9.5,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        _write_cell(
            deduct_table.cell(0, 2),
            rupiah_text(total_hm_denda),
            size=9.5,
            color=BA_RED if total_hm_denda < 0 else None,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )

        _write_cell(
            deduct_table.cell(1, 0),
            "Kompensasi :",
            size=9.5,
        )
        _write_cell(
            deduct_table.cell(1, 1),
            "",
            size=9.5,
        )
        _write_cell(
            deduct_table.cell(1, 2),
            rupiah_text(compensation),
            size=9.5,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )

        deduct_widths_cm = _autofit_table_columns(
            deduct_table,
            [
                [deduct_row0_texts[c], deduct_row1_texts[c]]
                for c in range(3)
            ],
        )

        _apply_thin_gray_borders(deduct_table)

        doc.add_paragraph()

        deduct_total_table = doc.add_table(rows=1, cols=3)
        deduct_total_table.style = "Table Grid"

        _write_cell(
            deduct_total_table.cell(0, 0),
            "Total Deduct :",
            bold=True,
            size=9.5,
        )
        _write_cell(
            deduct_total_table.cell(0, 1),
            "",
            size=9.5,
        )
        _write_cell(
            deduct_total_table.cell(0, 2),
            rupiah_text(total_deduct),
            bold=True,
            size=9.5,
            color=BA_RED if total_deduct < 0 else None,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            shading=BA_YELLOW,
        )

        _apply_thin_gray_borders(deduct_total_table)
        _set_row_double_border(deduct_total_table, 0)

        # Kolom disamakan persis lebarnya dengan deduct_table
        # di atas supaya kedua tabel terlihat sejajar rapi.
        deduct_total_table.autofit = False
        _set_table_fixed_layout(deduct_total_table)

        # Kolom 2 (nilai) perlu muat teks Total Deduct juga,
        # jadi dibandingkan dengan lebar hasil autofit-nya.
        deduct_total_widths_cm = list(
            deduct_widths_cm
        )
        deduct_total_widths_cm[2] = max(
            deduct_total_widths_cm[2],
            len(rupiah_text(total_deduct)) * 0.2
            + 0.45,
        )

        for row in deduct_total_table.rows:
            for idx, cell in enumerate(row.cells):
                cell.width = Cm(
                    deduct_total_widths_cm[idx]
                )

        if compensation_note:
            note_p = doc.add_paragraph()
            note_p.paragraph_format.space_before = Pt(6)
            note_run = note_p.add_run(
                f"Catatan Kompensasi: {compensation_note}"
            )
            note_run.italic = True
            note_run.font.size = Pt(9.5)

        doc.add_paragraph()

    # --------------------------------------------------------
    # CLOSING
    # --------------------------------------------------------

    closing = doc.add_paragraph()
    closing.add_run(
        "Demikian berita acara ini saya buat, Atas perhatian & "
        "kerjasamanya saya ucapkan terima kasih."
    ).font.size = Pt(10.5)

    place_date = doc.add_paragraph()
    place_date.add_run(
        f"{kota}, {tanggal_ba_text}"
    ).font.size = Pt(10.5)

    doc.add_paragraph()

    # --------------------------------------------------------
    # SIGNATURE BLOCK
    # --------------------------------------------------------

    mengetahui_list = [n for n in mengetahui_list if n.strip()]
    menyetujui_list = [n for n in menyetujui_list if n.strip()]

    if not mengetahui_list:
        mengetahui_list = [""]
    if not menyetujui_list:
        menyetujui_list = [""]

    total_cols = 1 + len(mengetahui_list) + len(menyetujui_list)

    sig_table = doc.add_table(rows=2, cols=total_cols)
    sig_table.autofit = True

    for row in sig_table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Header row: labels (merged per group)
    _write_cell(
        sig_table.cell(0, 0),
        "Dibuat oleh,",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10,
    )

    meng_start = 1
    meng_end = meng_start + len(mengetahui_list) - 1
    meng_cell = sig_table.cell(0, meng_start)
    for c in range(meng_start + 1, meng_end + 1):
        meng_cell = meng_cell.merge(sig_table.cell(0, c))
    _write_cell(
        meng_cell,
        "Mengetahui,",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10,
    )

    meny_start = meng_end + 1
    meny_end = meny_start + len(menyetujui_list) - 1
    meny_cell = sig_table.cell(0, meny_start)
    for c in range(meny_start + 1, meny_end + 1):
        meny_cell = meny_cell.merge(sig_table.cell(0, c))
    _write_cell(
        meny_cell,
        "Menyetujui,",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10,
    )

    # Spacer rows for signature space
    for _ in range(2):
        doc.add_paragraph()

    # Name row
    _write_cell(
        sig_table.cell(1, 0),
        f"( {dibuat_oleh} )" if dibuat_oleh else "(   )",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=10,
    )

    for i, name in enumerate(mengetahui_list):
        _write_cell(
            sig_table.cell(1, meng_start + i),
            f"( {name} )" if name else "(   )",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            size=10,
        )

    for i, name in enumerate(menyetujui_list):
        _write_cell(
            sig_table.cell(1, meny_start + i),
            f"( {name} )" if name else "(   )",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            size=10,
        )

    # Remove table borders for the signature block (it should
    # look like plain text columns, not a visible grid).
    tbl = sig_table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "none")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)
    tblPr.append(borders)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()


# ============================================================
# SIDEBAR
# ============================================================

st.title(
    "📦 Laporan Analisa Stock Opname"
)

st.caption(
    "Upload file DATA mentah untuk menghasilkan "
    "analisa otomatis dan sheet FINAL."
)

with st.sidebar:

    st.header(
        "⚙️ Parameter"
    )

    uploaded_file = st.file_uploader(
        "Input File",
        type=[
            "xlsx",
            "xls",
            "csv",
        ],
        help=(
            "Mendukung Excel XLSX/XLS dan CSV."
        ),
    )

    st.markdown("---")

    report_type = st.radio(
        "Jenis Laporan",
        [
            "SO IC - Sementara (Tim Audit)",
            "SO IC - Final (Tim Audit)",
            "SO Mandiri (Toko)",
        ],
        index=1,
        help=(
            "SO IC (Sementara / Final): lengkap dengan "
            "kolom HM/Total HM dan perhitungan deduct di "
            "Excel & Berita Acara. SO Mandiri: tanpa kolom "
            "HM/Total HM dan tanpa perhitungan deduct, "
            "untuk SO yang dilakukan store secara mandiri. "
            "Email: Sementara & Mandiri = tabel rekap; "
            "Final = lengkap dengan Total HM & perhitungan "
            "deduct."
        ),
    )

    # is_final = laporan lengkap (HM + deduct) => semua SO IC.
    is_final = report_type.startswith("SO IC")

    if report_type.startswith("SO IC - Sementara"):
        so_stage = "Sementara"
    elif report_type.startswith("SO IC - Final"):
        so_stage = "Final"
    else:
        so_stage = "Mandiri"

    if is_final:

        penalty_percent = st.number_input(
            "Persentase Denda (%)",
            min_value=0.0,
            max_value=100.0,
            value=30.0,
            step=1.0,
            format="%.2f",
        )

        compensation = st.number_input(
            "Nilai Kompensasi (Rp)",
            min_value=0.0,
            value=500_000.0,
            step=50_000.0,
            format="%.0f",
        )

        compensation_note = st.text_area(
            "Catatan Kompensasi",
            placeholder=(
                "Contoh: Kompensasi sesuai "
                "kebijakan periode berjalan."
            ),
        )

    else:

        penalty_percent = 0.0
        compensation = 0.0
        compensation_note = ""

    st.markdown("---")

    with st.expander(
        "📝 Data Berita Acara",
        expanded=False,
    ):

        ba_company = st.text_input(
            "Nama Perusahaan (Letterhead)",
            value="HARDWARE",
        )

        ba_kepada = st.text_input(
            "Kepada",
            value="",
        )

        ba_perihal = st.text_input(
            "Perihal",
            value=(
                "Posting Adjustment SO IC"
            ),
        )

        ba_no_adjustment = st.text_input(
            "No. ADJUSTMENT",
            value="",
            placeholder="ADJ/09/2026/000000001",
        )

        ba_tanggal_so = st.text_input(
            "Tanggal SO (untuk kalimat BA)",
            value="",
            placeholder="2 September 2026",
        )

        ba_kota = st.text_input(
            "Kota",
            value="Jakarta",
        )

        ba_tanggal_ba = st.text_input(
            "Tanggal Berita Acara",
            value="",
            placeholder="7 September 2026",
        )

        st.caption("Tanda tangan")

        ba_dibuat_oleh = st.text_input(
            "Dibuat oleh",
            value="",
        )

        ba_mengetahui = st.text_input(
            "Mengetahui (pisahkan koma)",
            value="",
            placeholder="Asenk, Bpk. Jeffry",
        )

        ba_menyetujui = st.text_input(
            "Menyetujui (pisahkan koma)",
            value="",
            placeholder="Ilham, Bpk. Yudi, Bpk. Andi Sutoyo",
        )

    with st.expander(
        "✉️ Data Email Analisa",
        expanded=False,
    ):

        email_sender = st.text_input(
            "Nama pengirim",
            value="Indra",
        )


# ============================================================
# NO FILE
# ============================================================

if uploaded_file is None:

    st.info(
        "Silakan upload file melalui sidebar."
    )

    st.markdown(
        """
### Kolom yang diperlukan

File mentah minimal harus memiliki:

- `OpnameNo`
- `Date`
- `Loc`
- `Golongan`
- `Artikel`
- `Barcode`
- `Qty POS`
- `Qty Fisik`
- `Price`

Kolom berikut **opsional** (kalau tidak ada, otomatis diisi 0):

- `HM`

Header dapat berada di bawah beberapa baris awal karena
aplikasi melakukan deteksi header secara otomatis.
"""
    )

    st.stop()


# ============================================================
# READ INPUT
# ============================================================

file_bytes = uploaded_file.getvalue()
filename = uploaded_file.name.lower()

try:

    if filename.endswith(".csv"):

        df_raw = read_csv_auto(
            file_bytes
        )

        selected_sheet = "CSV"

    else:

        if (
            filename.endswith(".xls")
            and not filename.endswith(
                ".xlsx"
            )
        ):
            extension = ".xls"
        else:
            extension = ".xlsx"

        excel = load_excel_file(
            file_bytes,
            extension,
        )

        sheets = excel.sheet_names

        # Prioritas DATA
        if "DATA" in sheets:
            default_index = sheets.index(
                "DATA"
            )

        elif "Data" in sheets:
            default_index = sheets.index(
                "Data"
            )

        else:
            default_index = 0

        selected_sheet = st.sidebar.selectbox(
            "Sheet Data",
            sheets,
            index=default_index,
        )

        df_raw = load_excel_sheet(
            excel,
            selected_sheet,
        )

except Exception as exc:

    st.error(
        f"Gagal membaca file: {exc}"
    )

    st.stop()


# ============================================================
# PROCESS
# ============================================================

try:

    hm_missing = (
        "HM"
        not in map_standard_columns(
            df_raw
        ).columns
    )

    df = prepare_data(
        df_raw
    )

    df = categorize(
        df
    )

    summary = create_summary(
        df
    )

    metadata = get_metadata(
        df
    )

    if hm_missing and is_final:
        st.warning(
            "⚠️ Kolom **HM** tidak ditemukan di file "
            "input. Nilai HM otomatis diisi 0 untuk "
            "seluruh baris."
        )

except Exception as exc:

    st.error(
        f"Gagal memproses data: {exc}"
    )

    with st.expander(
        "🔎 Header yang Terdeteksi"
    ):
        st.write(
            list(df_raw.columns)
        )

    st.stop()


# ============================================================
# DEDUCT CALCULATION
# ============================================================

total_hm = df[
    "Total HM"
].sum()

denda_value = (
    total_hm
    * penalty_percent
    / 100
)

total_hm_denda = (
    total_hm
    + denda_value
)

total_deduct = min(
    0,
    total_hm_denda
    + compensation,
)


# ============================================================
# METADATA
# ============================================================

metadata_line = (
    "📌 Metadata Laporan: "
    f"Lokasi {format_location_label(metadata['Loc']) or '-'} | "
    f"Tanggal SO {metadata['Date'] or '-'} | "
    f"No. Opname {metadata['OpnameNo'] or '-'}"
)

st.code(metadata_line, language=None)


# ============================================================
# KPI
# ============================================================

st.markdown(
    "## 📊 KPI"
)

grand = summary[
    summary["Kategori"]
    == "GRAND TOTAL"
].iloc[0]

if is_final:
    k1, k2, k3, k4, k5 = st.columns(5)
else:
    k1, k2, k3, k4 = st.columns(4)

with k1:
    st.metric(
        "Total Qty Fisik",
        f"{grand['Qty Fisik']:,.0f}",
    )

with k2:
    st.metric(
        "Total Qty POS",
        f"{grand['Qty POS']:,.0f}",
    )

with k3:
    st.metric(
        "Total Selisih Qty",
        f"{grand['Selisih']:,.0f}",
    )

with k4:
    st.metric(
        "Total Nilai Selisih",
        f"Rp {grand['Nilai Selisih (Rp)']:,.0f}",
    )

if is_final:
    with k5:
        st.metric(
            "Total Deduct",
            f"Rp {total_deduct:,.0f}",
        )


# ============================================================
# EXTRA KPI
# ============================================================

if is_final:

    a1, a2 = st.columns(2)

    with a1:
        st.metric(
            f"Total HM + Denda ({penalty_percent:.0f}%)",
            f"Rp {total_hm_denda:,.0f}",
        )

    with a2:
        st.metric(
            "Kompensasi",
            f"Rp {compensation:,.0f}",
        )


# ============================================================
# SUMMARY
# ============================================================

st.markdown(
    "## 📋 Rekapitulasi Kategori"
)

summary_display = summary.copy()

summary_numeric_cols = [
    "Qty Fisik",
    "Qty POS",
    "Selisih",
    "Nilai Selisih (Rp)",
]

if is_final:
    summary_numeric_cols += [
        "HM",
        "Total HM",
    ]
else:
    summary_display = summary_display.drop(
        columns=[
            "HM",
            "Total HM",
        ]
    )

for col in summary_numeric_cols:

    summary_display[col] = (
        summary_display[col]
        .map(
            lambda value:
            f"{value:,.0f}"
        )
    )

st.dataframe(
    summary_display,
    use_container_width=True,
    hide_index=True,
)


# ============================================================
# CATEGORY TABS
# ============================================================

with st.expander(
    "🔍 Detail Data Per Kategori",
    expanded=False,
):

    tabs = st.tabs(
        [
            (
                f"{category} "
                f"({len(df[df['Kategori'] == category])})"
            )
            for category in CATEGORIES
        ]
    )

    for tab, category in zip(
        tabs,
        CATEGORIES,
    ):

        with tab:

            part = df[
                df["Kategori"]
                == category
            ].copy()

            display_columns = [
                "Barcode",
                "Artikel",
                "Golongan",
                "Price",
                "Qty Fisik",
                "Qty POS",
                "Selisih",
                "Total",
                "Remark",
            ]

            column_config = {
                "Price":
                    st.column_config.NumberColumn(
                        "Price",
                        format="%,.0f",
                    ),

                "Qty Fisik":
                    st.column_config.NumberColumn(
                        "Qty Fisik",
                        format="%,.0f",
                    ),

                "Qty POS":
                    st.column_config.NumberColumn(
                        "Qty POS",
                        format="%,.0f",
                    ),

                "Selisih":
                    st.column_config.NumberColumn(
                        "Selisih",
                        format="%,.0f",
                    ),

                "Total":
                    st.column_config.NumberColumn(
                        "Total",
                        format="Rp %,.0f",
                    ),
            }

            if is_final:
                display_columns = (
                    display_columns[:8]
                    + [
                        "HM",
                        "Total HM",
                    ]
                    + display_columns[8:]
                )

                column_config[
                    "HM"
                ] = st.column_config.NumberColumn(
                    "HM",
                    format="%,.0f",
                )

                column_config[
                    "Total HM"
                ] = st.column_config.NumberColumn(
                    "Total HM",
                    format="Rp %,.0f",
                )

            st.dataframe(
                part[
                    display_columns
                ],
                use_container_width=True,
                hide_index=True,
                column_config=column_config,
            )

            if len(part) > 0:

                if is_final:
                    c1, c2, c3 = st.columns(3)
                else:
                    c1, c2 = st.columns(2)

                with c1:
                    st.metric(
                        "Qty Selisih",
                        f"{part['Selisih'].sum():,.0f}",
                    )

                with c2:
                    st.metric(
                        "Nilai Selisih",
                        f"Rp {part['Total'].sum():,.0f}",
                    )

                if is_final:
                    with c3:
                        st.metric(
                            "Total HM",
                            f"Rp {part['Total HM'].sum():,.0f}",
                        )


# ============================================================
# DEDUCT DETAIL
# ============================================================

st.markdown(
    "## 💰 Detail Perhitungan Deduct"
)

if is_final:

    d1, d2, d3, d4 = st.columns(4)

    with d1:
        st.metric(
            "Total HM",
            f"Rp {total_hm:,.0f}",
        )

    with d2:
        st.metric(
            f"Denda {penalty_percent:.0f}%",
            f"Rp {denda_value:,.0f}",
        )

    with d3:
        st.metric(
            "Kompensasi",
            f"Rp {compensation:,.0f}",
        )

    with d4:
        st.metric(
            "Total Deduct",
            f"Rp {total_deduct:,.0f}",
        )

    if compensation_note:
        st.info(
            f"Catatan Kompensasi: "
            f"{compensation_note}"
        )

else:

    st.caption(
        "Perhitungan deduct tidak berlaku untuk "
        "SO Mandiri."
    )


# ============================================================
# DATA PREVIEW
# ============================================================

with st.expander(
    "🔎 Preview Data Setelah Cleaning & Kategorisasi"
):

    preview_columns = [
        "OpnameNo",
        "Date",
        "Loc",
        "Golongan",
        "Artikel",
        "Barcode",
        "Qty POS",
        "Qty Fisik",
        "Price",
        "Selisih",
        "Total",
        "Kategori",
        "Remark",
    ]

    if is_final:
        preview_columns = (
            preview_columns[:9]
            + [
                "HM",
                "Total HM",
            ]
            + preview_columns[9:]
        )

    st.dataframe(
        df[
            preview_columns
        ],
        use_container_width=True,
        hide_index=True,
    )


# ============================================================
# EXPORT (AREA UNDUH LAPORAN)
# ============================================================

st.markdown("---")

st.markdown(
    "## 📦 Area Unduh Laporan"
)

# CSS: warna khusus tombol unduh Excel (hijau) & Word (biru),
# ditarget lewat key masing-masing tombol supaya tidak
# mempengaruhi tombol lain di halaman.
st.markdown(
    """
    <style>
    .st-key-download_excel_btn button {
        background-color: #1E8E3E;
        border-color: #1E8E3E;
        color: #FFFFFF;
    }
    .st-key-download_excel_btn button:hover {
        background-color: #167C33;
        border-color: #167C33;
        color: #FFFFFF;
    }
    .st-key-download_excel_btn button:focus {
        color: #FFFFFF;
    }
    .st-key-download_word_btn button {
        background-color: #1A73E8;
        border-color: #1A73E8;
        color: #FFFFFF;
    }
    .st-key-download_word_btn button:hover {
        background-color: #1558B0;
        border-color: #1558B0;
        color: #FFFFFF;
    }
    .st-key-download_word_btn button:focus {
        color: #FFFFFF;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

safe_loc = re.sub(
    r"[^A-Za-z0-9_-]+",
    "_",
    metadata["Loc"]
    or "ALL",
)

timestamp = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)

excel_bytes = None
ba_bytes = None

try:

    excel_bytes = create_final_excel(
        df=df,
        metadata=metadata,
        summary=summary,
        penalty_percent=(
            penalty_percent
            / 100
        ),
        compensation=compensation,
        compensation_note=(
            compensation_note
        ),
        include_hm=is_final,
        report_stage=so_stage,
    )

    file_prefix = (
        f"{so_stage.upper()}_STOCK_OPNAME_"
    )

    output_filename = (
        f"{file_prefix}"
        f"{safe_loc}_"
        f"{timestamp}.xlsx"
    )

except Exception as exc:

    st.error(
        f"Gagal membuat Excel "
        f"{so_stage.upper()}: {exc}"
    )

try:

    ba_bytes = create_berita_acara_docx(
        metadata=metadata,
        summary=summary,
        is_final=is_final,
        penalty_percent=penalty_percent,
        compensation=compensation,
        compensation_note=compensation_note,
        company_name=ba_company or "HARDWARE",
        kepada=ba_kepada,
        perihal=ba_perihal,
        no_adjustment=ba_no_adjustment,
        tanggal_so_text=(
            ba_tanggal_so
            or metadata["Date"]
            or "-"
        ),
        kota=ba_kota or "Jakarta",
        tanggal_ba_text=(
            ba_tanggal_ba
            or datetime.now().strftime("%d %B %Y")
        ),
        dibuat_oleh=ba_dibuat_oleh,
        mengetahui_list=[
            n.strip()
            for n in ba_mengetahui.split(",")
        ],
        menyetujui_list=[
            n.strip()
            for n in ba_menyetujui.split(",")
        ],
    )

    ba_filename = (
        f"BA_{so_stage.upper()}_SO_"
        f"{safe_loc}_{timestamp}.docx"
    )

except Exception as exc:

    st.error(
        f"Gagal membuat Berita Acara: {exc}"
    )

dl_col1, dl_col2 = st.columns(2)

with dl_col1:

    with st.container(border=True):

        st.markdown(
            "#### 📊 Laporan "
            + so_stage.upper()
        )

        st.caption(
            "Format: Microsoft Excel (.xlsx)"
        )

        if excel_bytes is not None:

            st.download_button(
                label="⬇️ Unduh Excel (.xlsx)",
                data=excel_bytes,
                file_name=output_filename,
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
                use_container_width=True,
                type="primary",
                key="download_excel_btn",
            )

with dl_col2:

    with st.container(border=True):

        st.markdown(
            "#### 📄 Berita Acara"
        )

        st.caption(
            "Format: Microsoft Word (.docx)"
        )

        if ba_bytes is not None:

            st.download_button(
                label="⬇️ Unduh Word (.docx)",
                data=ba_bytes,
                file_name=ba_filename,
                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                ),
                use_container_width=True,
                type="primary",
                key="download_word_btn",
            )


# ============================================================
# FORMAT EMAIL ANALISA SO
# ============================================================

st.markdown("---")

st.markdown(
    "## ✉️ Format Email Analisa SO"
)

email_subject, email_html = create_analisa_email(
    summary=summary,
    metadata=metadata,
    so_stage=so_stage,
    tanggal_so_text=ba_tanggal_so,
    sender_name=email_sender,
    penalty_percent=penalty_percent,
    compensation=compensation,
    compensation_note=compensation_note,
)

st.caption(
    "Subject (klik ikon salin di kanan atas kotak):"
)

st.code(
    email_subject,
    language=None,
)

st.caption(
    "Isi email — klik **Copy isi email**, lalu paste "
    "(Ctrl+V) di kolom compose Gmail/Outlook. Tabel, "
    "warna, dan format ikut tersalin."
)

render_email_preview(email_html)
